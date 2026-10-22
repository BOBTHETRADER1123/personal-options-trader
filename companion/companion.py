"""Outbound-only local Excel companion. No laptop port is opened.

Run with Python 3.11+: python companion.py
The local SQLite manifest makes a save followed by a lost acknowledgement safe.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import re
import secrets
import signal
import sqlite3
import tempfile
import threading
from pathlib import Path, PurePosixPath
from urllib.parse import urlsplit

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

log = logging.getLogger("paper-trading-companion")
REQUIRED_SHEETS = {"Trade Summary", "P&L Timeline", "Trade Events", "Strategy Legs", "Market Updates"}


def validate_backend_url(url: str) -> str:
    parsed = urlsplit(url)
    if not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise ValueError("Use the backend HTTPS origin, without credentials, query or fragment")
    if parsed.scheme != "https" and not (parsed.scheme == "http" and parsed.hostname in {"127.0.0.1", "localhost", "::1"}):
        raise ValueError("HTTPS is required except for a backend on this laptop")
    if parsed.path.strip("/"):
        raise ValueError("Use only the backend origin, without /api or another path")
    return url.rstrip("/")


def verify_file(path: Path, checksum: str, expected_size: int) -> None:
    if not path.is_file() or path.stat().st_size <= 0 or path.stat().st_size != expected_size:
        raise ValueError("Local Excel file is missing, empty, or has the wrong size")
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    if not secrets.compare_digest(digest.hexdigest(), checksum):
        raise ValueError("Local Excel checksum does not match the cloud manifest")
    # File streams also permit the temporary .partial suffix during validation.
    with path.open("rb") as source:
        book = load_workbook(source, read_only=True, keep_links=False)
        try:
            summary = dict(book["Trade Summary"].iter_rows(max_col=2, values_only=True)) if "Trade Summary" in book.sheetnames else {}
            required = {"Trade Summary", "Strategy Legs"} if summary.get("Report format") == "summary-v1" else REQUIRED_SHEETS
            if not required.issubset(book.sheetnames) or not summary.get("Trade ID"):
                raise ValueError("Local Excel workbook is missing required sheets")
            for sheet in book:
                for _ in sheet.iter_rows(values_only=True):
                    pass
        finally:
            book.close()


class LocalManifest:
    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(path)
        self.connection.execute("PRAGMA journal_mode=WAL")
        self.connection.execute("PRAGMA synchronous=FULL")
        self.connection.execute("CREATE TABLE IF NOT EXISTS jobs (id TEXT PRIMARY KEY, data TEXT NOT NULL, status TEXT NOT NULL, error TEXT)")
        self.connection.execute("CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        self.connection.commit()

    def identity(self) -> str:
        row = self.connection.execute("SELECT value FROM settings WHERE key='companion_id'").fetchone()
        if row:
            return row[0]
        identity = "laptop_" + secrets.token_hex(12)
        self.connection.execute("INSERT INTO settings (key,value) VALUES ('companion_id',?)", (identity,))
        self.connection.commit()
        return identity

    def put(self, job: dict, status: str, error: str | None = None) -> None:
        self.connection.execute(
            "INSERT INTO jobs(id,data,status,error) VALUES (?,?,?,?) ON CONFLICT(id) DO UPDATE SET data=excluded.data,status=excluded.status,error=excluded.error",
            (job["id"], json.dumps(job), status, error),
        )
        self.connection.commit()

    def pending(self) -> list[dict]:
        return [json.loads(row[0]) for row in self.connection.execute("SELECT data FROM jobs WHERE status != 'acknowledged'").fetchall()]

    def close(self) -> None:
        self.connection.close()


class Companion:
    def __init__(self, backend_url: str, token: str, folder: Path, client: httpx.Client | None = None):
        self.backend_url = validate_backend_url(backend_url)
        if len(token) < 32:
            raise ValueError("A paired companion token with at least 32 characters is required")
        self.folder = folder.expanduser().resolve()
        self.folder.mkdir(parents=True, exist_ok=True)
        self.manifest = LocalManifest(self.folder / ".companion" / "manifest.sqlite3")
        self.identity = self.manifest.identity()
        self.client = client or httpx.Client(base_url=self.backend_url, timeout=90, follow_redirects=False)
        self.client.headers.update({"Authorization": "Bearer " + token})

    def target(self, job: dict) -> Path:
        filename = job.get("filename", "")
        relative = PurePosixPath(str(job.get("relative_path", filename)).replace("\\", "/"))
        if not filename.endswith(".xlsx") or re.search(r'[<>:"/\\|?*\x00-\x1F]', filename):
            raise ValueError("Unsafe Excel filename")
        if relative.is_absolute() or any(part in {"..", "."} or ":" in part for part in relative.parts):
            raise ValueError("Unsafe report folder")
        if relative.name != filename:
            raise ValueError("Report path does not match the unique filename")
        target = self.folder.joinpath(*relative.parts).resolve()
        if self.folder not in target.parents:
            raise ValueError("Report path escapes the chosen folder")
        return target

    def request(self, method: str, path: str, job: dict | None = None, **kwargs) -> httpx.Response:
        headers = kwargs.pop("headers", {})
        if job and job.get("claim_token"):
            headers["X-Export-Claim"] = job["claim_token"]
        response = self.client.request(method, path, headers=headers, **kwargs)
        response.raise_for_status()
        return response

    def save_and_ack(self, job: dict) -> None:
        path = self.target(job)
        path.parent.mkdir(parents=True, exist_ok=True)
        if not path.exists():
            response = self.request("GET", f"/api/companion/jobs/{job['id']}/download", job)
            if not secrets.compare_digest(hashlib.sha256(response.content).hexdigest(), job["checksum"]):
                raise ValueError("Downloaded Excel checksum is incorrect")
            if len(response.content) != job["size"]:
                raise ValueError("Downloaded Excel size is incorrect")
            temp_path = None
            try:
                with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".partial", delete=False) as stream:
                    temp_path = Path(stream.name)
                    stream.write(response.content)
                    stream.flush()
                    os.fsync(stream.fileno())
                verify_file(temp_path, job["checksum"], job["size"])
                # Never overwrite an unrelated or changed workbook. Another
                # companion may have completed this path while we downloaded.
                if path.exists():
                    verify_file(path, job["checksum"], job["size"])
                else:
                    try:
                        # Linking is an atomic create that cannot overwrite a
                        # file appearing between the existence check and save.
                        os.link(temp_path, path)
                    except FileExistsError:
                        verify_file(path, job["checksum"], job["size"])
                    except OSError:
                        # Filesystems without hard links use exclusive create.
                        # A crash here cannot trigger cloud cleanup: reopening
                        # and final checksum validation precede acknowledgement.
                        with path.open("xb") as target, temp_path.open("rb") as source:
                            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                                target.write(chunk)
                            target.flush()
                            os.fsync(target.fileno())
            finally:
                if temp_path:
                    temp_path.unlink(missing_ok=True)
        verify_file(path, job["checksum"], job["size"])
        job["local_path"] = str(path)
        self.manifest.put(job, "verified_pending_ack")
        acknowledgement = {
            "companion_id": self.identity, "checksum": job["checksum"], "size": path.stat().st_size,
            "filename": job["filename"], "local_path": str(path), "workbook_valid": True,
        }
        self.request("POST", f"/api/companion/jobs/{job['id']}/ack", job, json=acknowledgement)
        self.manifest.put(job, "acknowledged")
        log.info("Saved and verified: %s", path)

    def _handle(self, job: dict) -> None:
        try:
            self.save_and_ack(job)
        except Exception as exc:
            # Keep the claim in the local durable manifest. If the server was
            # restarted after local saving, the next pass re-verifies and acks.
            self.manifest.put(job, "retry", type(exc).__name__)
            log.warning("Job %s awaits retry (%s). Cloud records are retained.", job.get("id"), type(exc).__name__)
            if not isinstance(exc, httpx.HTTPError):
                try:
                    self.request("POST", f"/api/companion/jobs/{job['id']}/failed", job,
                                 json={"message": f"File could not be saved/verified ({type(exc).__name__})."})
                except httpx.HTTPError:
                    pass

    def poll_once(self) -> None:
        for pending in self.manifest.pending():
            self._handle(pending)
        response = self.request("GET", "/api/companion/jobs")
        for candidate in response.json().get("jobs", []):
            try:
                claim = self.request("POST", f"/api/companion/jobs/{candidate['id']}/claim",
                                     json={"companion_id": self.identity}).json()
            except httpx.HTTPStatusError as exc:
                if exc.response.status_code == 409:
                    continue
                raise
            if claim.get("already_saved"):
                continue
            self.manifest.put(claim, "claimed")
            self._handle(claim)

    def close(self) -> None:
        self.client.close()
        self.manifest.close()


def main() -> None:
    load_dotenv(Path(__file__).with_name(".env"))
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    stop = threading.Event()
    signal.signal(signal.SIGINT, lambda *_: stop.set())
    signal.signal(signal.SIGTERM, lambda *_: stop.set())
    backend = os.getenv("BACKEND_URL", "http://127.0.0.1:8000")
    folder = Path(os.getenv("TRADES_FOLDER", str(Path.home() / "Documents" / "PaperTrading" / "Trades")))
    service = Companion(backend, os.getenv("COMPANION_TOKEN", ""), folder)
    interval = max(5, int(os.getenv("POLL_SECONDS", "10")))
    log.info("Companion is watching for completed reports. Save folder: %s", folder)
    try:
        while not stop.is_set():
            try:
                service.poll_once()
            except (httpx.HTTPError, OSError) as exc:
                log.warning("Backend or folder unavailable (%s); jobs remain queued.", type(exc).__name__)
            stop.wait(interval)
    finally:
        service.close()


if __name__ == "__main__":
    main()
