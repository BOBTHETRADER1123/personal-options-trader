"""Durable, recoverable trade-report jobs. The database journal is authoritative.

The report payload and XLSX bytes are stored in the same database as the trade;
the optional filesystem copy is only a convenience, never the only cloud copy.
No workbook write or network request happens inside the closing-trade transaction.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import logging
import os
import re
import tempfile
import time
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import JSON, Boolean, Integer, LargeBinary, String, Text, func, or_, select, update
from sqlalchemy.orm import Mapped, Session, mapped_column

from .database import Base, SessionLocal

log = logging.getLogger(__name__)
IST = ZoneInfo("Asia/Kolkata")
REQUIRED_SHEETS = {"Trade Summary", "P&L Timeline", "Trade Events", "Strategy Legs", "Market Updates"}
MAX_EXCEL_ROWS = 1_048_576
_delete_enabled: Callable[[Session], bool] = lambda db: False
_delete_handler: Callable[[Session, str], dict[str, int]] | None = None


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def parse_time(value: Any) -> datetime:
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, (int, float)):
        result = datetime.fromtimestamp(value / (1000 if value > 1e12 else 1), timezone.utc)
    else:
        result = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return result.replace(tzinfo=timezone.utc) if result.tzinfo is None else result


def local_excel_time(value: Any) -> Any:
    """Excel cannot retain tzinfo: keep local datetime plus an explicit IST label."""
    if value in (None, ""):
        return None
    try:
        return parse_time(value).astimezone(IST).replace(tzinfo=None)
    except (ValueError, TypeError, OverflowError):
        return safe_cell(value)


def safe_cell(value: Any) -> Any:
    # Quotes, trade names and symbols can originate outside the app. Never turn
    # them into Excel formula expressions (including malicious external links).
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
        return value[:32767]
    if isinstance(value, float) and (value != value or abs(value) == float("inf")):
        return "Unavailable / unbounded"
    return value


def safe_name(value: Any, fallback: str = "trade") -> str:
    result = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value)).strip("._")[:100]
    return result or fallback


class ExportJob(Base):
    __tablename__ = "export_jobs"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    trade_id: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    broker: Mapped[str] = mapped_column(String(40))
    filename: Mapped[str] = mapped_column(String(255))
    relative_path: Mapped[str] = mapped_column(String(500))
    status: Mapped[str] = mapped_column(String(24), default="pending", index=True)
    payload: Mapped[dict] = mapped_column(JSON)
    file_bytes: Mapped[bytes | None] = mapped_column(LargeBinary, nullable=True)
    checksum: Mapped[str | None] = mapped_column(String(64), nullable=True)
    size: Mapped[int] = mapped_column(Integer, default=0)
    attempts: Mapped[int] = mapped_column(Integer, default=0)
    error: Mapped[str | None] = mapped_column(Text, nullable=True)
    created_at: Mapped[str] = mapped_column(String(40), default=now_iso)
    updated_at: Mapped[str] = mapped_column(String(40), default=now_iso)
    generated_at: Mapped[str | None] = mapped_column(String(40), nullable=True)
    saved_at: Mapped[str | None] = mapped_column(String(40), nullable=True)
    lease_until: Mapped[str | None] = mapped_column(String(40), nullable=True)
    claim_hash: Mapped[str | None] = mapped_column(String(64), nullable=True)
    companion_id: Mapped[str | None] = mapped_column(String(80), nullable=True)
    local_path: Mapped[str | None] = mapped_column(Text, nullable=True)
    verification_status: Mapped[str] = mapped_column(String(40), default="not_verified")
    deletion_status: Mapped[str] = mapped_column(String(40), default="retained")
    deleted_at: Mapped[str | None] = mapped_column(String(40), nullable=True)
    deleted_counts: Mapped[dict | None] = mapped_column(JSON, nullable=True)
    automatic_retry: Mapped[bool] = mapped_column(Boolean, default=True)


class DailyExportReconciliation(Base):
    __tablename__ = "export_daily_reconciliation"
    date: Mapped[str] = mapped_column(String(10), primary_key=True)
    generated_at: Mapped[str] = mapped_column(String(40), default=now_iso)
    data: Mapped[dict] = mapped_column(JSON)


def configure_archival(
    enabled_checker: Callable[[Session], bool],
    delete_handler: Callable[[Session, str], dict[str, int]] | None = None,
) -> None:
    """The caller must use an explicit owner-controlled setting, default false.

    A handler may delete trade-owned rows only. Shared quotes, credentials,
    active states and the permanent export manifest must remain untouched.
    """
    global _delete_enabled, _delete_handler
    _delete_enabled, _delete_handler = enabled_checker, delete_handler


def enqueue_export(
    db: Session,
    trade: dict[str, Any],
    events: list[dict[str, Any]],
    timeline: list[dict[str, Any]],
    quotes: list[dict[str, Any]],
) -> ExportJob:
    if str(trade.get("status", "")).lower() not in {"closed", "cancelled"}:
        raise ValueError("Only completed trades can be archived")
    trade_id = str(trade["id"])
    existing = db.scalar(select(ExportJob).where(ExportJob.trade_id == trade_id))
    if existing:
        return existing
    closed = parse_time(trade.get("closed_at") or now_iso()).astimezone(IST)
    filename = f"{safe_name(trade.get('underlying', 'OPTIONS'))}_{safe_name(trade_id)}_{closed:%Y%m%d_%H%M%S_%f}_IST.xlsx"
    broker = str(trade.get("broker", "Unknown"))
    job = ExportJob(
        id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"personal-options-trader:export:{trade_id}")),
        trade_id=trade_id,
        broker=broker,
        filename=filename,
        relative_path=f"{safe_name(broker)}/{closed:%Y/%m}/Trade_{safe_name(trade_id)}/{filename}",
        status="pending",
        payload={"trade": trade, "events": events, "timeline": timeline, "quotes": quotes},
        created_at=now_iso(),
        updated_at=now_iso(),
    )
    db.add(job)
    db.flush()
    return job


def public_job(job: ExportJob) -> dict[str, Any]:
    result = {
        key: getattr(job, key)
        for key in (
            "id", "trade_id", "broker", "filename", "relative_path", "status", "checksum", "size",
            "attempts", "error", "created_at", "updated_at", "generated_at", "saved_at",
            "verification_status", "deletion_status", "deleted_at", "deleted_counts",
        )
    }
    result["demo"] = bool(job.payload.get("trade", {}).get("demo"))
    return result


def _flatten(record: dict[str, Any]) -> dict[str, Any]:
    inner = record.get("data")
    return {**record, **inner} if isinstance(inner, dict) else record


def _table(wb: Workbook, title: str, columns: list[tuple[str, str]], records: list[dict]) -> None:
    """Split oversized logs into more sheets rather than silently dropping rows."""
    chunks = [records[i:i + MAX_EXCEL_ROWS - 2] for i in range(0, len(records), MAX_EXCEL_ROWS - 2)] or [[]]
    for chunk_index, chunk in enumerate(chunks):
        sheet_title = title if not chunk_index else f"{title[:25]} {chunk_index + 1}"
        ws = wb.create_sheet(sheet_title)
        ws.append([label for _, label in columns])
        for raw in chunk:
            record = _flatten(raw)
            values = []
            for key, label in columns:
                value = record if key == "data" else record.get(key)
                if key in {"received_ns", "receive_ns"} and value is not None:
                    value = str(value)  # Excel numbers retain only 15 significant digits.
                if key.endswith("_at") or key in {"timestamp", "event_timestamp", "exchange_timestamp", "received_at"}:
                    value = local_excel_time(value)
                else:
                    value = safe_cell(value)
                values.append(value)
            ws.append(values)
        _style_table(ws)


def _style_table(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="152638")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(40, max(18, len(str(ws.cell(1, col).value)) + 3))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color="162435")
            cell.alignment = Alignment(vertical="top")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5F8")
            if isinstance(cell.value, datetime):
                cell.number_format = "dd mmm yyyy hh:mm:ss.000"
                ws.column_dimensions[cell.column_letter].width = 28
            elif isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'


def build_workbook(payload: dict[str, Any]) -> bytes:
    trade = payload["trade"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Trade Summary"
    ws.sheet_view.showGridLines = False
    ws.append(["PERSONAL OPTIONS TRADER", "COMPLETED PAPER TRADE"])
    ws.append(["Time zone", "Asia/Kolkata (IST, UTC+05:30)"])
    ws.append(["Mode", "PAPER ONLY — all fills and P&L are simulated"])
    ws.append(["Record source", "Persistent application event journal; quotes retain their individual broker source"])
    ws.append(["Metric", "Value"])
    rows = [
        ("Trade ID", trade.get("id")), ("Strategy ID", trade.get("strategy_id")),
        ("Strategy name", trade.get("name")), ("Broker", trade.get("broker")),
        ("Underlying", trade.get("underlying")), ("Status", trade.get("status")),
        ("Quote and contract origin", "SYNTHETIC SAMPLE — fictional contracts and simulated quotes" if trade.get("demo") else "Broker master and received market data; orders remain simulated"),
        ("Entry time (IST)", local_excel_time(trade.get("created_at"))),
        ("Closing time (IST)", local_excel_time(trade.get("closed_at"))),
        ("Gross P&L (INR)", trade.get("gross_pnl")),
        ("Estimated brokerage and charges (INR)", trade.get("charges")),
        ("Net closing P&L (INR)", trade.get("net_pnl")),
        ("Realized gross P&L (INR)", trade.get("realized_pnl")),
        ("Unrealized gross P&L (INR)", trade.get("unrealized_pnl")),
        ("Net premium cash flow (INR)", trade.get("net_premium")),
        ("Maximum favourable excursion / MFE (INR)", trade.get("mfe")),
        ("MFE time (IST)", local_excel_time(trade.get("mfe_at"))),
        ("Maximum adverse excursion / MAE (INR)", trade.get("mae")),
        ("MAE time (IST)", local_excel_time(trade.get("mae_at"))),
        ("Profit giveback (INR)", max(0, float(trade.get("mfe") or 0) - float(trade.get("net_pnl") or 0))),
    ]
    risk = trade.get("risk") or {}
    for field, label in [
        ("max_profit", "Theoretical maximum profit (INR)"), ("max_loss", "Theoretical maximum loss (INR)"),
        ("margin", "Estimated margin (INR)"), ("estimated_margin", "Overall estimated margin (INR)"),
        ("hedge_benefit", "Estimated hedge benefit (INR)"), ("breakevens", "Expiry breakevens"),
        ("assumptions", "Risk assumptions"),
    ]:
        rows.append((label, risk.get(field)))
    rows.extend([
        ("Accounting method", "Weighted-average entry; realized P&L on closed quantity; charges tracked separately."),
        ("Recording note", "Market Updates contains received events only. Carried-forward one-second snapshots are explicitly identified."),
        ("Precision note", "Receive timestamps are runtime timestamps; precision does not establish exchange accuracy or delivery latency."),
    ])
    for name, value in rows:
        ws.append([name, safe_cell(value)])
    ws.column_dimensions["A"].width = 49
    ws.column_dimensions["B"].width = 96
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="152638")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=14)
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5F8")
            if isinstance(cell.value, datetime):
                cell.number_format = "dd mmm yyyy hh:mm:ss.000"
            elif isinstance(cell.value, (float, int)):
                cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
    ws.freeze_panes = "B6"

    timeline_rows = []
    leg_timeline_rows = []
    for raw in payload.get("timeline", []):
        item = _flatten(raw)
        legs = item.get("legs", [])
        if legs:
            item = {
                **item, "instrument": trade.get("underlying"),
                "open_lots": sum(float(leg.get("open_lots", 0)) for leg in legs),
            }
            if len(legs) == 1:
                item = {**item, **{key: legs[0].get(key) for key in ("ltp", "bid", "ask", "feed_age", "event_id")},
                        "entry_price": legs[0].get("average_entry_price")}
            for leg in legs:
                leg_timeline_rows.append({**leg, "timestamp": item.get("timestamp"), "broker": item.get("broker", trade.get("broker")),
                                          "carried_forward": item.get("carried_forward"), "entry_price": leg.get("average_entry_price")})
        timeline_rows.append(item)
    _table(wb, "P&L Timeline", [
        ("timestamp", "Snapshot / event time (IST)"), ("broker", "Broker"),
        ("instrument", "Instrument"), ("ltp", "LTP"), ("bid", "Bid"), ("ask", "Ask"),
        ("open_lots", "Open lots"), ("entry_price", "Average entry price"),
        ("unrealized_pnl", "Unrealized gross P&L (INR)"), ("realized_pnl", "Realized gross P&L (INR)"),
        ("charges", "Estimated charges (INR)"), ("net_pnl", "Net paper P&L (INR)"),
        ("feed_age", "Feed age (seconds)"), ("carried_forward", "Carried-forward snapshot"),
        ("event_id", "Event / sequence ID"), ("kind", "Record type"), ("data", "Complete source record"),
    ], timeline_rows)
    if leg_timeline_rows:
        _table(wb, "Leg Timeline", [
            ("timestamp", "Snapshot / event time (IST)"), ("broker", "Broker"),
            ("instrument_id", "Instrument ID"), ("symbol", "Trading symbol"),
            ("ltp", "LTP"), ("bid", "Bid"), ("ask", "Ask"), ("open_lots", "Open lots"),
            ("entry_price", "Average entry price"), ("unrealized_pnl", "Unrealized gross P&L (INR)"),
            ("realized_pnl", "Realized gross P&L (INR)"), ("charges", "Estimated charges (INR)"),
            ("net_pnl", "Net leg paper P&L (INR)"), ("feed_age", "Feed age (seconds)"),
            ("carried_forward", "Carried-forward snapshot"), ("event_id", "Event / sequence ID"),
            ("data", "Complete source record"),
        ], leg_timeline_rows)
    _table(wb, "Trade Events", [
        ("timestamp", "Event time (IST)"), ("kind", "Event type"), ("id", "Journal ID"),
        ("action_id", "Action ID"), ("leg_id", "Leg ID"), ("lots", "Lots"),
        ("price", "Simulated price"), ("data", "Complete event details"),
    ], payload.get("events", []))
    _table(wb, "Strategy Legs", [
        ("id", "Leg ID"), ("instrument_id", "Instrument ID"), ("symbol", "Trading symbol"),
        ("underlying", "Underlying"), ("exchange", "Exchange"), ("expiry", "Expiry"),
        ("strike", "Strike"), ("option_type", "CE / PE"), ("side", "Buy / sell"),
        ("lots", "Total entry lots"), ("open_lots", "Remaining lots"), ("lot_size", "Contract lot size"),
        ("quantity", "Total contract quantity"), ("price_multiplier", "Monetary price multiplier"),
        ("monetary_quantity", "Quantity used for P&L"), ("multiplier_verified", "Multiplier verified from master"),
        ("entry_price", "Average entry price"),
        ("exit_price", "Average exit price"), ("current_price", "Last valuation price"),
        ("entry_at", "Entry time (IST)"), ("exit_at", "Exit time (IST)"),
        ("realized_pnl", "Realized gross P&L (INR)"), ("unrealized_pnl", "Unrealized gross P&L (INR)"),
        ("charges", "Estimated charges (INR)"), ("status", "Leg status"),
        ("order_type", "Simulated order type"),
    ], [{**leg,
         "quantity": float(leg.get("lots", 0)) * float(leg.get("lot_size", 0)),
         "price_multiplier": float(leg.get("price_multiplier") or 1),
         "monetary_quantity": float(leg.get("lots", 0)) * float(leg.get("lot_size", 0)) * float(leg.get("price_multiplier") or 1),
         } for leg in trade.get("legs", [])])
    _table(wb, "Market Updates", [
        ("received_at", "Local receive time (IST)"), ("received_ns", "Receive Unix nanoseconds"),
        ("exchange_timestamp", "Exchange time (IST, if supplied)"), ("broker", "Actual source broker"),
        ("instrument_id", "Instrument ID"), ("exchange", "Exchange"), ("segment", "Segment"),
        ("symbol", "Trading symbol"), ("expiry", "Expiry"), ("strike", "Strike"),
        ("option_type", "CE / PE"), ("ltp", "LTP"), ("bid", "Bid"), ("ask", "Ask"),
        ("volume", "Volume"), ("open_interest", "Open interest"), ("sequence", "Provider sequence"),
        ("id", "Journal ID"), ("data", "Complete source packet"),
    ], payload.get("quotes", []))
    # Plot existing captured data only. A sparse timeline stays sparse.
    timeline_sheet = wb["P&L Timeline"]
    if 2 <= timeline_sheet.max_row <= 10000:
        chart = LineChart()
        chart.title = "Recorded net paper P&L"
        chart.y_axis.title = "INR"
        chart.x_axis.title = "Recorded snapshots, in order"
        chart.add_data(Reference(timeline_sheet, min_col=12, min_row=1, max_row=timeline_sheet.max_row), titles_from_data=True)
        chart.height, chart.width = 9, 22
        ws.add_chart(chart, f"A{ws.max_row + 3}")
    stream = io.BytesIO()
    wb.save(stream)
    value = stream.getvalue()
    verify_workbook(value)
    return value


def verify_workbook(content: bytes) -> None:
    if not content:
        raise ValueError("Excel report is empty")
    with io.BytesIO(content) as stream:
        wb = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            if not REQUIRED_SHEETS.issubset(wb.sheetnames):
                raise ValueError("Excel report is missing required sheets")
            for sheet in wb:
                if not sheet.max_row:
                    raise ValueError("Excel report has an empty sheet")
                # Consume every XML row: merely opening a ZIP is insufficient.
                for _ in sheet.iter_rows(values_only=True):
                    pass
        finally:
            wb.close()


def _optional_disk_copy(job: ExportJob) -> None:
    directory = os.getenv("EXPORT_ROOT", "").strip()
    if not directory or not job.file_bytes:
        return
    root = Path(directory).expanduser().resolve()
    target = (root / job.relative_path).resolve()
    if root not in target.parents:
        raise ValueError("Unsafe export path")
    target.parent.mkdir(parents=True, exist_ok=True)
    temp_name = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".tmp", delete=False) as tmp:
            temp_name = Path(tmp.name)
            tmp.write(job.file_bytes)
            tmp.flush()
            os.fsync(tmp.fileno())
        os.replace(temp_name, target)
    finally:
        if temp_name is not None:
            temp_name.unlink(missing_ok=True)


def process_pending_once(session_factory=SessionLocal, limit: int = 3) -> int:
    """CAS claims survive multiple callers; expired leases recover after restart."""
    processed = 0
    with session_factory() as db:
        cutoff = now_iso()
        candidates = db.scalars(select(ExportJob.id).where(or_(
            ExportJob.status == "pending",
            (ExportJob.status == "generating") & (ExportJob.lease_until < cutoff),
            (ExportJob.status == "failed") & (ExportJob.automatic_retry == True) & (ExportJob.lease_until < cutoff),
        )).order_by(ExportJob.created_at).limit(limit)).all()
    for job_id in candidates:
        with session_factory() as db:
            current = now_iso()
            lease = (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat(timespec="microseconds")
            claimed = db.execute(update(ExportJob).where(
                ExportJob.id == job_id,
                or_(ExportJob.status == "pending", (ExportJob.status.in_(["failed", "generating"])) & (ExportJob.lease_until < current)),
            ).values(status="generating", lease_until=lease, updated_at=current, attempts=ExportJob.attempts + 1))
            db.commit()
            if claimed.rowcount != 1:
                continue
            job = db.get(ExportJob, job_id)
            try:
                content = build_workbook(job.payload)
                job.file_bytes = content
                job.checksum = hashlib.sha256(content).hexdigest()
                job.size = len(content)
                job.generated_at = now_iso()
                job.updated_at = now_iso()
                job.status = "ready"
                job.lease_until = None
                job.error = None
                # A locked optional convenience copy must not lose the durable
                # database archive; the companion can still download the bytes.
                try:
                    _optional_disk_copy(job)
                except OSError:
                    log.warning("Optional report folder unavailable for job %s; database copy retained", job.id)
                db.commit()
                processed += 1
            except Exception as exc:
                db.rollback()
                job = db.get(ExportJob, job_id)
                job.status = "failed"
                job.error = f"Workbook generation failed ({type(exc).__name__}); source journal retained."
                job.updated_at = now_iso()
                seconds = min(3600, 15 * (2 ** min(job.attempts, 8)))
                job.lease_until = (datetime.now(timezone.utc) + timedelta(seconds=seconds)).isoformat(timespec="microseconds")
                db.commit()
                log.exception("Workbook generation failed for job %s", job_id)
    return processed


async def reporting_loop(stop_event: asyncio.Event | None = None) -> None:
    stop_event = stop_event or asyncio.Event()
    last_reconciled = 0.0
    while not stop_event.is_set():
        try:
            await asyncio.to_thread(process_pending_once)
            if time.monotonic() - last_reconciled > 60:
                await asyncio.to_thread(save_daily_reconciliation)
                last_reconciled = time.monotonic()
        except Exception:
            log.exception("Report worker iteration failed; durable jobs remain queued")
        try:
            await asyncio.wait_for(stop_event.wait(), timeout=2)
        except asyncio.TimeoutError:
            pass


def archive_after_verified_save(db: Session, job: ExportJob) -> None:
    """Run only inside the acknowledgement transaction; no broad deletion."""
    if job.status != "saved" or job.verification_status != "verified" or not job.checksum or not job.saved_at:
        raise ValueError("Local Excel verification is required before archival")
    if job.deletion_status == "archived":
        return
    if not _delete_enabled(db):
        job.deletion_status = "retained"
        return
    if _delete_handler is None:
        job.deletion_status = "retained_no_archive_handler"
        return
    counts = _delete_handler(db, job.trade_id)
    job.deletion_status = "archived"
    job.deleted_at = now_iso()
    job.deleted_counts = counts
    # Keep the complete verified XLSX and a small permanent manifest. This does
    # not remove shared quotes, active trades, settings, or any other trade.
    job.payload = {"trade": job.payload.get("trade", {}), "archived": True}


def reconciliation(db: Session) -> dict[str, Any]:
    """Each call is a daily reconciliation view; no invented successful saves."""
    jobs = db.scalars(select(ExportJob).order_by(ExportJob.created_at.desc())).all()
    counts: dict[str, int] = {}
    for job in jobs:
        counts[job.status] = counts.get(job.status, 0) + 1
    return {
        "date": datetime.now(IST).date().isoformat(),
        "timezone": "Asia/Kolkata", "generated_at": now_iso(), "total": len(jobs), "counts": counts,
        "pending": sum(job.status in {"pending", "generating", "ready", "claimed"} for job in jobs),
        "saved": counts.get("saved", 0), "failed": counts.get("failed", 0),
        "mismatches": [public_job(job) for job in jobs if job.status == "saved" and job.verification_status != "verified"],
        "jobs": [public_job(job) for job in jobs],
    }


def save_daily_reconciliation(session_factory=SessionLocal) -> dict[str, Any]:
    with session_factory() as db:
        report = reconciliation(db)
        row = db.get(DailyExportReconciliation, report["date"])
        if row is None:
            db.add(DailyExportReconciliation(date=report["date"], generated_at=report["generated_at"], data=report))
        else:
            row.generated_at, row.data = report["generated_at"], report
        db.commit()
        return report
