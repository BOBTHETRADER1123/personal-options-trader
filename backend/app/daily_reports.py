"""Durable IST daily exports: one worksheet per independent broker trade.
Only captured PnlSnapshot history is exported; missing seconds are never invented.
"""
from __future__ import annotations
import asyncio
import csv
import gzip
import hashlib
import io
import json
import logging
import os
import time
import threading
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import Float, Integer, JSON, LargeBinary, String, func, select, update
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Mapped, mapped_column
from .database import Base, SessionLocal, iso_now, run_transaction
from .models import PnlSnapshot, RecordingGap, Trade
from .reporting import parse_time, safe_cell
from .scheduler import ExchangeCalendar

IST=ZoneInfo('Asia/Kolkata');UTC=timezone.utc
log=logging.getLogger(__name__)
XLSX_MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
CSV_MIME='application/gzip'
COLUMNS=(
 'timestamp_utc','timestamp_ist','broker','trade_id','strategy_id','strategy_name','underlying',
 'leg_id','contract_identifier','exchange','segment','symbol','instrument_type','strike','expiry',
 'buy_sell','lots','contract_lot_size','total_quantity','price_multiplier','entry_at','entry_price',
 'exit_at','exit_price','holding_seconds','latest_price','quote_received_at','exchange_quote_at',
 'quote_age_seconds','quote_reused','stale','data_freshness','leg_unrealized_pnl','leg_realized_pnl',
 'leg_gross_pnl','leg_net_pnl','leg_charges','total_unrealized_pnl','total_realized_pnl','gross_pnl',
 'estimated_charges','net_pnl','observed_max_profit_lifecycle','observed_max_loss_lifecycle',
 'estimated_initial_margin','estimated_hedge_benefit','estimated_final_margin','net_option_premium',
 'theoretical_max_profit','theoretical_max_loss','breakevens','leg_signed_quantity',
 'strategy_signed_notional_estimate','close_reason','backup_status_at_export','session_status','demo')
META_ROWS=8

class DailyReport(Base):
 __tablename__='daily_reports'
 trading_date:Mapped[str]=mapped_column(String(10),primary_key=True)
 status:Mapped[str]=mapped_column(String(24),default='pending',index=True)
 trade_count:Mapped[int]=mapped_column(Integer,default=0)
 strategy_count:Mapped[int]=mapped_column(Integer,default=0)
 snapshot_count:Mapped[int]=mapped_column(Integer,default=0)
 record_count:Mapped[int]=mapped_column(Integer,default=0)
 worksheet_count:Mapped[int]=mapped_column(Integer,default=0)
 total_size:Mapped[int]=mapped_column(Integer,default=0)
 manifest_checksum:Mapped[str|None]=mapped_column(String(64),nullable=True)
 manifest:Mapped[dict]=mapped_column(JSON,default=dict)
 source_signature:Mapped[str|None]=mapped_column(String(64),nullable=True)
 attempts:Mapped[int]=mapped_column(Integer,default=0)
 lease_until:Mapped[float]=mapped_column(Float,default=0)
 lease_owner:Mapped[str|None]=mapped_column(String(64),nullable=True)
 next_attempt:Mapped[float]=mapped_column(Float,default=0)
 error:Mapped[str|None]=mapped_column(String(300),nullable=True)
 created_at:Mapped[str]=mapped_column(String(40),default=iso_now)
 finalized_at:Mapped[str|None]=mapped_column(String(40),nullable=True)
 local_ack_checksum:Mapped[str|None]=mapped_column(String(64),nullable=True)
 local_ack_at:Mapped[str|None]=mapped_column(String(40),nullable=True)

class DailyReportPart(Base):
 __tablename__='daily_report_parts'
 id:Mapped[str]=mapped_column(String(80),primary_key=True)
 trading_date:Mapped[str]=mapped_column(String(10),index=True)
 kind:Mapped[str]=mapped_column(String(20))
 part_number:Mapped[int]=mapped_column(Integer)
 filename:Mapped[str]=mapped_column(String(180))
 mime_type:Mapped[str]=mapped_column(String(100))
 size:Mapped[int]=mapped_column(Integer)
 sha256:Mapped[str]=mapped_column(String(64))
 md5:Mapped[str]=mapped_column(String(32))
 record_count:Mapped[int]=mapped_column(Integer)
 worksheet_count:Mapped[int]=mapped_column(Integer)
 metadata_json:Mapped[dict]=mapped_column(JSON)
 file_bytes:Mapped[bytes]=mapped_column(LargeBinary,deferred=True)

class DailyReportError(Exception):pass

@dataclass(frozen=True)
class ReportLimits:
 max_sheet_rows:int=1_048_576
 max_workbook_cells:int=8_000_000
 max_workbook_bytes:int=64*1024*1024
 max_data_rows_per_sheet:int=100_000
 lease_seconds:int=1800
 @classmethod
 def from_env(cls):
  return cls(max_workbook_cells=max(10000,min(8_000_000,int(os.getenv('DAILY_EXPORT_MAX_CELLS','8000000')))),
   max_workbook_bytes=max(4096,int(os.getenv('DAILY_EXPORT_MAX_FILE_MB','64'))*1024*1024),
   max_data_rows_per_sheet=max(1,int(os.getenv('DAILY_EXPORT_SHEET_DATA_ROWS','100000'))))
 def validate(self):
  if not META_ROWS<self.max_sheet_rows<=1_048_576:raise DailyReportError('Invalid Excel worksheet row limit.')
  if self.max_workbook_cells<len(COLUMNS)*(META_ROWS+1)+100:raise DailyReportError('Cell allowance cannot hold a trade worksheet.')
  if self.max_data_rows_per_sheet<1 or self.max_workbook_bytes<4096:raise DailyReportError('Invalid daily report limits.')

def bounds(trading_date):
 day=date.fromisoformat(trading_date)
 if day.isoformat()!=trading_date:raise DailyReportError('Use an ISO trading date.')
 start=datetime.combine(day,datetime.min.time(),IST)
 return int(start.timestamp()),int((start+timedelta(days=1)).timestamp())

def _canonical(value):return json.dumps(value,sort_keys=True,separators=(',',':'),ensure_ascii=False,default=str).encode()

def ensure_day(db,trading_date):
 bounds(trading_date)
 report=db.get(DailyReport,trading_date)
 if report:return report
 report=DailyReport(trading_date=trading_date,status='pending',manifest={},lease_until=0)
 db.add(report);db.flush();return report

def queue_day(db,trading_date):
 report=ensure_day(db,trading_date)
 if report.status!='finalized' and report.lease_until<=time.time():report.status,report.error,report.next_attempt='pending',None,0
 db.commit();return public_report(report)

def public_report(report):
 return {key:getattr(report,key) for key in ('trading_date','status','trade_count','strategy_count','snapshot_count','record_count',
  'worksheet_count','total_size','manifest_checksum','attempts','error','created_at','finalized_at','local_ack_checksum','local_ack_at')}

def list_reports(db,limit=90):
 return [public_report(row) for row in db.scalars(select(DailyReport).order_by(DailyReport.trading_date.desc()).limit(max(1,min(366,limit))))]

def report_manifest(db,trading_date):
 report=db.get(DailyReport,trading_date)
 if not report:raise DailyReportError('Daily report not found.')
 return {**public_report(report),'manifest':report.manifest}

def artifact(db,trading_date,kind='xlsx',part=1):
 report=db.get(DailyReport,trading_date)
 if not report or report.status!='finalized':raise DailyReportError('Daily export is not finalized; source records retained.')
 row=db.get(DailyReportPart,f'{trading_date}:{kind}:{int(part):04d}')
 if not row:raise DailyReportError('Daily report part not found.')
 payload=row.file_bytes
 if len(payload)!=row.size or hashlib.sha256(payload).hexdigest()!=row.sha256:raise DailyReportError('Daily export integrity verification failed.')
 return {**row.metadata_json,'file_bytes':payload}

def cloud_artifacts(db,trading_date):
 """Yield verified bytes one part at a time, never all daily binary content."""
 report=db.get(DailyReport,trading_date)
 if not report or report.status!='finalized':return
 for item in report.manifest.get('artifacts',[]):yield artifact(db,trading_date,item['kind'],item['part_number'])

def verify_local_ack(db,trading_date,manifest_checksum,confirmed,confirmed_record_count):
 report=db.get(DailyReport,trading_date)
 if not report or report.status!='finalized' or not confirmed:raise DailyReportError('Explicit confirmation of a verified finalized local backup is required.')
 if report.manifest_checksum!=manifest_checksum or confirmed_record_count!=report.record_count:raise DailyReportError('Checksum or record count does not match the finalized daily export.')
 report.local_ack_checksum,report.local_ack_at=manifest_checksum,iso_now();db.commit();return public_report(report)

def _day_trades(db,trading_date):
 start,end=bounds(trading_date)
 ids=select(PnlSnapshot.trade_id).where(PnlSnapshot.snapshot_second>=start,PnlSnapshot.snapshot_second<end).distinct()
 selected={row.id:row.data for row in db.scalars(select(Trade).where(Trade.id.in_(ids)))}
 for row in db.scalars(select(Trade)):
  t=row.data
  try:overlaps=bool(t.get('created_at') and parse_time(t['created_at']).timestamp()<end and (not t.get('closed_at') or parse_time(t['closed_at']).timestamp()>=start))
  except (ValueError,TypeError,OverflowError):overlaps=False
  if overlaps:selected[row.id]=t
 return sorted(selected.values(),key=lambda t:(t.get('created_at',''),t['id']))

def _source_signature(db,trading_date,trades):
 start,end=bounds(trading_date)
 rows=db.execute(select(PnlSnapshot.trade_id,func.count(),func.min(PnlSnapshot.snapshot_second),func.max(PnlSnapshot.snapshot_second)).where(
  PnlSnapshot.snapshot_second>=start,PnlSnapshot.snapshot_second<end).group_by(PnlSnapshot.trade_id)).all()
 return hashlib.sha256(_canonical({'date':trading_date,'trades':sorted(t['id'] for t in trades),'snapshots':sorted([list(r) for r in rows])})).hexdigest()

def _finalization_allowed(trading_date,trades,calendar,now):
 day=date.fromisoformat(trading_date)
 if day>now.astimezone(IST).date():return False,'This trading date has not begun.'
 # Synthetic practice recording runs outside exchange hours, including weekends.
 # Keep its current IST day mutable until midnight so later seconds are included.
 if any(t.get('demo') for t in trades) and now.timestamp()<bounds(trading_date)[1]:
  return False,'Synthetic practice records finalize after the IST date ends.'
 ends=[]
 for t in trades:
  if t.get('demo'):continue
  for leg in t.get('legs',[]):
   intervals=calendar.intervals(leg,day)
   if intervals is None:return False,'No verified exchange calendar for a contract on this date.'
   ends.extend(end for _,end in intervals)
 if not trades:return now.timestamp()>=bounds(trading_date)[1],'An empty day finalizes after the IST date ends.'
 if ends and now.timestamp()<max(ends):return False,'The relevant verified trading sessions have not ended.'
 return True,'Verified sessions ended.'

def _duration(entry,end):
 try:return max(0,int((parse_time(end)-parse_time(entry)).total_seconds())) if entry and end else None
 except (ValueError,TypeError,OverflowError):return None

def _row(t,s,l,day_end):
 original=next((v for v in t.get('legs',[]) if v.get('id')==l.get('id')),{})
 at=s.get('timestamp_utc') or s.get('timestamp');risk=s.get('risk') or {}
 lots=l.get('open_lots',l.get('lots',0)) or 0;qty=lots*(l.get('lot_size') or 0)
 entry_at=original.get('entry_at') or t.get('created_at');exit_at=original.get('exit_at')
 try:
  if exit_at and parse_time(exit_at).timestamp()>=day_end:exit_at=None
 except (ValueError,TypeError):exit_at=None
 exposure=sum((1 if item.get('side')=='BUY' else -1)*(item.get('open_lots') or 0)*(item.get('lot_size') or 0)*
  (item.get('price_multiplier') or 1)*(item.get('latest_price') or 0) for item in s.get('legs',[]))
 values=[at,s.get('timestamp_ist') or (parse_time(at).astimezone(IST).isoformat(timespec='seconds') if at else None),
  t.get('broker'),t['id'],t.get('strategy_id') or t['id'],t.get('name'),t.get('underlying'),l.get('id'),l.get('instrument_id'),l.get('exchange'),
  l.get('segment'),l.get('symbol'),l.get('option_type'),l.get('strike'),l.get('expiry'),l.get('side'),lots,l.get('lot_size'),qty,
  l.get('price_multiplier') or 1,entry_at,l.get('entry_price'),exit_at,original.get('exit_price') if exit_at else None,_duration(entry_at,exit_at or at),
  l.get('latest_price',l.get('ltp')),l.get('quote_received_at',l.get('quote_at')),l.get('exchange_quote_at'),l.get('quote_age_seconds'),l.get('quote_reused'),
  l.get('stale'),l.get('data_freshness'),l.get('unrealized_pnl'),l.get('realized_pnl'),l.get('gross_pnl'),l.get('net_pnl'),l.get('charges'),
  s.get('unrealized_pnl'),s.get('realized_pnl'),s.get('gross_pnl'),s.get('charges'),s.get('net_pnl'),s.get('mfe'),s.get('mae'),
  risk.get('estimated_initial_margin'),s.get('hedge_benefit',risk.get('hedge_benefit')),s.get('margin_estimate',risk.get('estimated_margin')),s.get('net_premium'),
  'Unlimited' if risk.get('profit_unlimited') else risk.get('max_profit'),'Unlimited' if risk.get('loss_unlimited') else risk.get('max_loss'),
  risk.get('breakevens'),qty*(1 if l.get('side')=='BUY' else -1),exposure,t.get('close_reason') if exit_at else None,
  'Generated locally; remote verification is tracked in app',l.get('session'),bool(t.get('demo'))]
 return [safe_cell(v) for v in values]

def _normalize_xlsx(payload):
 output=io.BytesIO()
 with zipfile.ZipFile(io.BytesIO(payload)) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
  for name in sorted(source.namelist()):
   item=zipfile.ZipInfo(name,date_time=(2000,1,1,0,0,0));item.compress_type=zipfile.ZIP_DEFLATED;target.writestr(item,source.read(name))
 return output.getvalue()

def _workbook(trading_date):
 wb=Workbook(write_only=True);fixed=datetime.combine(date.fromisoformat(trading_date),datetime.min.time())
 wb.properties.created=wb.properties.modified=fixed;return wb

def _heading(sheet,values):
 cells=[]
 for v in values:
  cell=WriteOnlyCell(sheet,value=safe_cell(v));cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='15394C');cells.append(cell)
 sheet.append(cells)

def _sheet_name(ordinal,broker,continued,chunk):
 label={'kotak':'Kotak','fivepaisa':'5paisa'}.get(broker,'Legacy')
 suffix=('_Continued' if continued else '')+(f'_{chunk}' if chunk>1 else '')
 return f'Trade_{ordinal:03d}_{label}{suffix}'[:31]

def _part_metadata(day,kind,part,content,records,sheets,coverage):
 extension='xlsx' if kind in ('xlsx','index_xlsx') else 'csv.gz'
 suffix='Index' if kind=='index_xlsx' else f'Part_{part:03d}'
 return {'id':f'{day}:{kind}:{part:04d}','trading_date':day,'kind':kind,'part_number':part,
  'filename':f'Paper_Trading_{day}_{suffix}.{extension}','mime_type':XLSX_MIME if kind in ('xlsx','index_xlsx') else CSV_MIME,
  'size':len(content),'sha256':hashlib.sha256(content).hexdigest(),'md5':hashlib.md5(content).hexdigest(),
  'record_count':records,'worksheet_count':sheets,'coverage':coverage}

class _FileTooLarge(DailyReportError):pass

class DailyTradeIndex(Base):
 __tablename__='daily_trade_indices'
 trade_id:Mapped[str]=mapped_column(String(80),primary_key=True)
 ordinal:Mapped[int]=mapped_column(Integer,unique=True)

class DailyTradeCounter(Base):
 __tablename__='daily_trade_counter'
 id:Mapped[int]=mapped_column(Integer,primary_key=True)
 value:Mapped[int]=mapped_column(Integer,default=0)

class DailyReportService:
 def __init__(self,session_factory=SessionLocal,calendar=None,limits=None,on_finalized=None):
  self.factory=session_factory;self.calendar=calendar or ExchangeCalendar();self.limits=limits or ReportLimits.from_env()
  self.limits.validate();self.last_error=None;self.on_finalized=on_finalized
  self._wake_lock=threading.Lock();self._loop=None;self._wake=None;self._pending_wake=False

 def notify_work(self):
  with self._wake_lock:
   loop,wake=self._loop,self._wake
   if loop is None or wake is None:self._pending_wake=True;return
  if not loop.is_closed():
   try:loop.call_soon_threadsafe(wake.set)
   except RuntimeError:pass

 def discover_dates(self):
  with self.factory() as db:
   first,last=db.execute(select(func.min(PnlSnapshot.snapshot_second),func.max(PnlSnapshot.snapshot_second))).one()
   if first is None:return 0
   day=datetime.fromtimestamp(first,UTC).astimezone(IST).date();last_day=datetime.fromtimestamp(last,UTC).astimezone(IST).date();count=0
   while day<=last_day:
    key=day.isoformat();start,end=bounds(key)
    exists=db.scalar(select(PnlSnapshot.trade_id).where(PnlSnapshot.snapshot_second>=start,PnlSnapshot.snapshot_second<end).limit(1))
    if exists and not db.get(DailyReport,key):ensure_day(db,key);count+=1
    day+=timedelta(days=1)
   db.commit();return count

 def _claim(self,trading_date,now):
  owner=uuid.uuid4().hex
  def claim(db):
   report=ensure_day(db,trading_date)
   if report.status=='finalized':return None
   trades=_day_trades(db,trading_date);allowed,reason=_finalization_allowed(trading_date,trades,self.calendar,now)
   if not allowed:report.error=reason;return None
   current=time.time()
   changed=db.execute(update(DailyReport).where(DailyReport.trading_date==trading_date,DailyReport.status!='finalized',
    DailyReport.lease_until<=current,DailyReport.next_attempt<=current).values(status='generating',lease_owner=owner,lease_until=current+self.limits.lease_seconds,
    attempts=DailyReport.attempts+1,error=None))
   return owner if changed.rowcount==1 else None
  return run_transaction(claim,self.factory)

 def _ordinal_map(self,trades):
  def assign(db):
   if db.get_bind().dialect.name=='postgresql':from sqlalchemy.dialects.postgresql import insert
   else:from sqlalchemy.dialects.sqlite import insert
   # This first statement obtains SQLite's writer lock before any reads.
   db.execute(insert(DailyTradeCounter).values(id=1,value=0).on_conflict_do_nothing(index_elements=['id']))
   result={}
   for t in trades:
    row=db.get(DailyTradeIndex,t['id'])
    if not row:
     ordinal=db.scalar(update(DailyTradeCounter).where(DailyTradeCounter.id==1).values(value=DailyTradeCounter.value+1).returning(DailyTradeCounter.value))
     row=DailyTradeIndex(trade_id=t['id'],ordinal=ordinal);db.add(row);db.flush()
    result[t['id']]=row.ordinal
   return result
  return run_transaction(assign,self.factory)

 def _snapshots(self,trade_id,start,end):
  cursor=start-1
  while True:
   # Release the read transaction before writing a completed export part.
   # This works with rollback-journal SQLite as well as verified WAL storage.
   with self.factory() as db:
    rows=list(db.execute(select(PnlSnapshot.snapshot_second,PnlSnapshot.data).where(PnlSnapshot.trade_id==trade_id,
     PnlSnapshot.snapshot_second>cursor,PnlSnapshot.snapshot_second<end).order_by(PnlSnapshot.snapshot_second).limit(250)))
   if not rows:return
   for second,payload in rows:yield second,payload
   cursor=rows[-1][0]

 def _store(self,owner,metadata,payload):
  def store(db):
   report=db.get(DailyReport,metadata['trading_date'])
   if not report or report.status!='generating' or report.lease_owner!=owner or report.lease_until<=time.time():
    raise DailyReportError('Daily export ownership changed; source records retained.')
   values={key:metadata[key] for key in ('id','trading_date','kind','part_number','filename','mime_type','size','sha256','md5','record_count','worksheet_count')}
   existing=db.get(DailyReportPart,metadata['id'])
   if existing:
    # Unpublished staging artifacts may be rebuilt after an interrupted attempt.
    # Finalized archives cannot enter this branch or be modified.
    if existing.sha256!=metadata['sha256']:
     for key,value in values.items():setattr(existing,key,value)
     existing.metadata_json=metadata;existing.file_bytes=payload
   else:db.add(DailyReportPart(**values,metadata_json=metadata,file_bytes=payload))
  run_transaction(store,self.factory)

 def _finish_part(self,wb,summary,csv_text,day,number,coverage,owner):
  records=sum(item['record_count'] for item in coverage)
  summary.append(['Personal Options Trader','Daily export',day,'Asia/Kolkata'])
  summary.append(['Broker-specific trades in this part',len({item['trade_id'] for item in coverage})])
  summary.append(['Leg-level one-second records in this part',records])
  summary.append(['Snapshots are observations; reused quotes retain original timestamps.'])
  summary.append(['Estimated margins/charges are separate from official broker funds and theoretical payoff.'])
  _heading(summary,['Trade ID','Strategy ID','Broker','Worksheet','Records','Snapshot count','Continued'])
  for item in coverage:summary.append([safe_cell(item[key]) for key in ('trade_id','strategy_id','broker','sheet','record_count','snapshot_count','continued')])
  stream=io.BytesIO();wb.save(stream);xlsx=_normalize_xlsx(stream.getvalue())
  csv_bytes=gzip.compress(csv_text.getvalue().encode('utf-8-sig'),mtime=0)
  if max(len(xlsx),len(csv_bytes))>self.limits.max_workbook_bytes:raise _FileTooLarge('Generated daily file exceeds the configured byte limit.')
  verified=load_workbook(io.BytesIO(xlsx),read_only=True,data_only=True)
  try:
   if set(verified.sheetnames)!={'Daily_Summary',*(item['sheet'] for item in coverage)}:raise DailyReportError('Trade-sheet coverage verification failed.')
   for item in coverage:
    count=sum(1 for _ in verified[item['sheet']].iter_rows(values_only=True))-META_ROWS
    if count!=item['record_count']:raise DailyReportError('Second-by-second worksheet row verification failed.')
  finally:verified.close()
  count=sum(1 for _ in csv.reader(io.StringIO(gzip.decompress(csv_bytes).decode('utf-8-sig'))))-1
  if count!=records:raise DailyReportError('Compressed history row count verification failed.')
  artifacts=[]
  for kind,content,sheets in (('xlsx',xlsx,len(coverage)+1),('csv_gz',csv_bytes,0)):
   metadata=_part_metadata(day,kind,number,content,records,sheets,coverage);self._store(owner,metadata,content);artifacts.append(metadata)
  return artifacts

 def _build(self,day,owner,cell_limit):
  with self.factory() as db:
   trades=_day_trades(db,day);signature=_source_signature(db,day,trades)
  start,end=bounds(day);ordinals=self._ordinal_map(trades)
  artifacts=[];all_coverage=[];snapshots_total=0
  def fresh():
   book=_workbook(day);summary=book.create_sheet('Daily_Summary');text=io.StringIO(newline='');writer=csv.writer(text);writer.writerow(COLUMNS)
   return book,summary,text,writer
  wb,summary,csv_text,csv_writer=fresh();coverage=[];cells=400;part_number=1
  max_rows=min(self.limits.max_data_rows_per_sheet,self.limits.max_sheet_rows-META_ROWS,max(1,(cell_limit-400)//len(COLUMNS)-META_ROWS))
  for t in trades:
   continued=parse_time(t.get('created_at',datetime.fromtimestamp(start,UTC))).astimezone(IST).date().isoformat()<day
   chunk=0;sheet=None;current=None
   for second,payload in self._snapshots(t['id'],start,end):
    snapshots_total+=1
    for leg_index,leg in enumerate(payload.get('legs') or [{}]):
     if sheet is None or current['record_count']>=max_rows or cells+len(COLUMNS)>cell_limit:
      if coverage and cells+len(COLUMNS)*(META_ROWS+1)>cell_limit:
       artifacts.extend(self._finish_part(wb,summary,csv_text,day,part_number,coverage,owner));all_coverage.extend(coverage)
       part_number+=1;wb,summary,csv_text,csv_writer=fresh();coverage=[];cells=400
      chunk+=1;name=_sheet_name(ordinals[t['id']],t.get('broker'),continued,chunk);sheet=wb.create_sheet(name)
      current={'trade_id':t['id'],'strategy_id':t.get('strategy_id') or t['id'],'broker':t.get('broker'),'sheet':name,
       'record_count':0,'snapshot_count':0,'continued':continued,'part_number':part_number};coverage.append(current)
      for values in (['Trading date',day,'Asia/Kolkata'],['Trade ID',t['id']],['Strategy ID',t.get('strategy_id') or t['id']],
       ['Broker',t.get('broker')],['Lifecycle','Continued' if continued else 'Started today'],
       ['Recording','Only captured observations; gaps are not reconstructed.'],['Backup','Verify manifest and every part before cleanup.']):
       sheet.append([safe_cell(value) for value in values])
      _heading(sheet,COLUMNS);sheet.freeze_panes=f'A{META_ROWS+1}';cells+=len(COLUMNS)*META_ROWS
     row=_row(t,payload,leg,end);sheet.append(row);csv_writer.writerow(row);current['record_count']+=1
     if leg_index==0:current['snapshot_count']+=1
     cells+=len(COLUMNS)
   if sheet is None:
    if coverage and cells+len(COLUMNS)*META_ROWS>cell_limit:
     artifacts.extend(self._finish_part(wb,summary,csv_text,day,part_number,coverage,owner));all_coverage.extend(coverage)
     part_number+=1;wb,summary,csv_text,csv_writer=fresh();coverage=[];cells=400
    name=_sheet_name(ordinals[t['id']],t.get('broker'),continued,1);sheet=wb.create_sheet(name)
    for values in (['Trading date',day],['Trade ID',t['id']],['Strategy ID',t.get('strategy_id') or t['id']],['Broker',t.get('broker')],
     ['Snapshot count',0],['Data status','No one-second observations captured on this date.'],['Important','No rows were invented.']):sheet.append(values)
    _heading(sheet,COLUMNS)
    coverage.append({'trade_id':t['id'],'strategy_id':t.get('strategy_id') or t['id'],'broker':t.get('broker'),'sheet':name,
     'record_count':0,'snapshot_count':0,'continued':continued,'part_number':part_number});cells+=len(COLUMNS)*META_ROWS
  artifacts.extend(self._finish_part(wb,summary,csv_text,day,part_number,coverage,owner));all_coverage.extend(coverage)
  with self.factory() as db:
   if _source_signature(db,day,_day_trades(db,day))!=signature:raise DailyReportError('Source records changed during export; not finalized.')
   gap_rows=db.execute(select(RecordingGap.trade_id,RecordingGap.start_second,RecordingGap.end_second,RecordingGap.reason).where(
    RecordingGap.start_second<end,RecordingGap.end_second>=start)).all()
  gaps=[{'trade_id':t,'start_second':max(lo,start),'end_second':min(hi,end-1),'reason':reason} for t,lo,hi,reason in gap_rows]
  if part_number>1:
   index=_workbook(day);table=index.create_sheet('Daily_Summary')
   _heading(table,['Daily master index',day,'Download and verify every listed file.'])
   _heading(table,['Kind','Part','Filename','Records','Size bytes','SHA-256'])
   for item in artifacts:table.append([item[key] for key in ('kind','part_number','filename','record_count','size','sha256')])
   _heading(table,['Trade ID','Strategy ID','Broker','Worksheet','Part','Records','Continued'])
   for item in all_coverage:table.append([safe_cell(item[key]) for key in ('trade_id','strategy_id','broker','sheet','part_number','record_count','continued')])
   stream=io.BytesIO();index.save(stream);content=_normalize_xlsx(stream.getvalue())
   if len(content)>self.limits.max_workbook_bytes:raise DailyReportError('Master index exceeds configured file limit; records retained.')
   metadata=_part_metadata(day,'index_xlsx',0,content,0,1,[]);self._store(owner,metadata,content);artifacts.append(metadata)
  record_count=sum(item['record_count'] for item in all_coverage)
  manifest={'format':'personal-options-daily-v1','trading_date':day,'timezone':'Asia/Kolkata','trade_count':len(trades),
   'strategy_count':len({t.get('strategy_id') or t['id'] for t in trades}),'snapshot_count':snapshots_total,'record_count':record_count,
   'coverage':all_coverage,'artifacts':artifacts,'recording_gaps':gaps,'source_signature':signature,
   'records_definition':'One captured leg per trade-second; snapshot_count counts distinct trade-seconds.',
   'all_source_records_retained':True,'data_freshness':'Original quote timestamps retained; gaps never reconstructed.',
   'warnings':['Margin, costs and hedge benefit are estimates.','Theoretical expiry payoff differs from observed P&L.',
    'A local download is not proof of an independent backup. Verify every file and retain an additional copy.']}
  checksum=hashlib.sha256(_canonical(manifest)).hexdigest()
  def finalize(db):
   report=db.get(DailyReport,day)
   if report.lease_owner!=owner or report.status!='generating' or report.lease_until<=time.time():raise DailyReportError('Export ownership changed before finalization.')
   report.status='finalized';report.finalized_at=iso_now();report.lease_until=0;report.lease_owner=None;report.error=None
   report.trade_count=len(trades);report.strategy_count=manifest['strategy_count'];report.snapshot_count=snapshots_total;report.record_count=record_count
   report.worksheet_count=sum(a['worksheet_count'] for a in artifacts);report.total_size=sum(a['size'] for a in artifacts)
   report.manifest=manifest;report.manifest_checksum=checksum;report.source_signature=signature;return public_report(report)
  result=run_transaction(finalize,self.factory)
  if self.on_finalized:self.on_finalized()
  return result

 def generate_day(self,trading_date,now=None):
  moment=now or datetime.now(UTC)
  if moment.tzinfo is None:raise DailyReportError('Daily finalization requires timezone-aware current time.')
  owner=self._claim(trading_date,moment)
  if owner is None:
   with self.factory() as db:return report_manifest(db,trading_date)
  try:
   cells=self.limits.max_workbook_cells
   while True:
    try:return self._build(trading_date,owner,cells)
    except _FileTooLarge:
     minimum=len(COLUMNS)*(META_ROWS+1)+400
     if cells<=minimum:raise DailyReportError('Even a single-record workbook exceeds the configured byte limit; source records retained.')
     cells=max(minimum,cells//2)
  except Exception as error:
   safe=str(error) if isinstance(error,DailyReportError) else f'Daily export interrupted ({type(error).__name__}); source records retained.'
   def fail(db):
    db.execute(update(DailyReport).where(DailyReport.trading_date==trading_date,DailyReport.lease_owner==owner).values(
     status='failed',error=safe[:300],lease_until=0,lease_owner=None,next_attempt=time.time()+min(3600,30*2**min(7,db.get(DailyReport,trading_date).attempts))))
   run_transaction(fail,self.factory);self.last_error=safe
   if isinstance(error,DailyReportError):raise
   raise DailyReportError(safe) from None

 async def run(self,stop_event):
  loop=asyncio.get_running_loop();wake=asyncio.Event()
  with self._wake_lock:
   self._loop,self._wake=loop,wake
   if self._pending_wake:wake.set();self._pending_wake=False
  try:await asyncio.to_thread(self.discover_dates)
  except SQLAlchemyError:log.warning('Daily date recovery deferred; source records retained.')
  while not stop_event.is_set():
   try:
    with self.factory() as db:
     days=list(db.scalars(select(DailyReport.trading_date).where(DailyReport.status.in_(['pending','failed','generating']),
      DailyReport.lease_until<=time.time(),DailyReport.next_attempt<=time.time()).order_by(DailyReport.trading_date).limit(7)))
    for day in days:
     if stop_event.is_set():break
     try:await asyncio.to_thread(self.generate_day,day)
     except DailyReportError:log.warning('Daily export deferred; source records retained.')
   except SQLAlchemyError:log.warning('Daily export database unavailable; retry deferred.')
   shutdown=asyncio.create_task(stop_event.wait());changed=asyncio.create_task(wake.wait())
   try:await asyncio.wait((shutdown,changed),timeout=60,return_when=asyncio.FIRST_COMPLETED)
   finally:
    for task in (shutdown,changed):task.cancel()
    await asyncio.gather(shutdown,changed,return_exceptions=True)
    wake.clear()
