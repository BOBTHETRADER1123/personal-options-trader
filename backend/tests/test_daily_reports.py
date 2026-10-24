"""Daily exports use isolated SQLite and mocked official calendars only."""
from dataclasses import replace
from datetime import datetime,timedelta,timezone
import csv
import gzip
import hashlib
import io
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine,select,func,update
from sqlalchemy.orm import sessionmaker
from app.database import Base,configure_sqlite
from app.models import Trade,PnlSnapshot,RecordingGap
from app.scheduler import ExchangeCalendar
from app import daily_reports as daily

UTC=timezone.utc
DAY='2026-10-26'
OPEN=datetime(2026,10,26,3,45,tzinfo=UTC)
AFTER=datetime(2026,10,26,10,11,tzinfo=UTC)

@pytest.fixture
def factory(tmp_path):
 sql=create_engine(f'sqlite:///{tmp_path / "daily.db"}')
 configure_sqlite(sql);Base.metadata.create_all(sql)
 yield sessionmaker(sql,expire_on_commit=False)
 sql.dispose()

@pytest.fixture
def calendar():
 return ExchangeCalendar({'rules':[{'exchange':'NSE','segments':['nse_fo'],'valid_from':'2026-01-01','valid_through':'2026-12-31',
  'sources':['https://exchange.example/official-mocked-test-calendar'],'verified_at':'2026-10-24','sessions':[['09:15','15:40']],
  'weekdays':[0,1,2,3,4],'holidays':[]}]})

def insert_trade(db,identifier='t001',broker='kotak',strategy='s001',day=OPEN,seconds=3,legs=2,created=None,closed=None):
 data={'id':identifier,'broker':broker,'strategy_id':strategy,'name':'=UNTRUSTED_FORMULA()','underlying':'NIFTY','status':'open',
  'created_at':(created or day).isoformat(),'closed_at':closed.isoformat() if closed else None,'demo':False,
  'risk':{'max_profit':None,'profit_unlimited':True,'max_loss':1000,'estimated_initial_margin':2000,'estimated_margin':1000,'hedge_benefit':1000},
  'legs':[{'id':f'leg{i}','instrument_id':f'{broker}:nse_fo:{i}','exchange':'NSE','segment':'nse_fo','underlying':'NIFTY',
   'expiry':'2026-11-03','option_type':'CE','side':'BUY' if i%2==0 else 'SELL','lots':1,'open_lots':1,'lot_size':50,'entry_price':100,
   'entry_at':(created or day).isoformat()} for i in range(legs)]}
 db.add(Trade(id=identifier,broker=broker,status='open',data=data))
 for n in range(seconds):add_snapshot(db,data,day+timedelta(seconds=n))
 db.commit();return data

def add_snapshot(db,trade,at):
 payload={'timestamp':at.isoformat(),'timestamp_utc':at.isoformat(),'timestamp_ist':at.astimezone(daily.IST).isoformat(),
  'trade_id':trade['id'],'broker':trade['broker'],'strategy_id':trade['strategy_id'],'net_pnl':10,'charges':5,'gross_pnl':15,
  'unrealized_pnl':15,'realized_pnl':0,'mfe':20,'mae':-10,'risk':trade['risk'],
  'legs':[{**leg,'latest_price':101,'quote_at':(at-timedelta(seconds=2)).isoformat(),'quote_age_seconds':2,'quote_reused':True,
   'stale':False,'data_freshness':'reused','unrealized_pnl':15,'realized_pnl':0,'charges':5,'session':{'known':True,'open':True}}
   for leg in trade['legs']]}
 db.add(PnlSnapshot(trade_id=trade['id'],snapshot_second=int(at.timestamp()),timestamp=at.isoformat(),data=payload))


def make(factory,calendar,**limits):return daily.DailyReportService(factory,calendar,replace(daily.ReportLimits(),**limits))

def test_realistic_20_strategies_mirrored_40_trade_sheets_and_complete_rows(factory,calendar):
 with factory() as db:
  for n in range(20):
   for broker in ('kotak','fivepaisa'):insert_trade(db,f't{n:03d}_{broker}',broker,f's{n:03d}')
 result=make(factory,calendar).generate_day(DAY,AFTER)
 assert result['status']=='finalized'
 assert (result['trade_count'],result['strategy_count'],result['snapshot_count'],result['record_count'])==(40,20,120,240)
 with factory() as db:
  manifest=daily.report_manifest(db,DAY)['manifest'];item=daily.artifact(db,DAY)
  assert len(manifest['coverage'])==40 and len(manifest['artifacts'])==2
  wb=load_workbook(io.BytesIO(item['file_bytes']),read_only=True,data_only=True)
  assert len(wb.sheetnames)==41 and 'Daily_Summary' in wb.sheetnames
  for covered in manifest['coverage']:
   assert sum(1 for _ in wb[covered['sheet']].iter_rows(values_only=True))==daily.META_ROWS+6
  wb.close()
  csv_item=daily.artifact(db,DAY,'csv_gz');rows=list(csv.DictReader(io.StringIO(gzip.decompress(csv_item['file_bytes']).decode('utf-8-sig'))))
  assert len(rows)==240 and len({r['trade_id'] for r in rows})==40
  assert rows[0]['quote_reused']=='True' and rows[0]['quote_age_seconds']=='2'
  assert rows[0]['strategy_name'].startswith("'=")
  assert db.scalar(select(func.count()).select_from(PnlSnapshot))==120


def test_current_session_and_unknown_calendar_do_not_finalize(factory,calendar):
 with factory() as db:insert_trade(db)
 service=make(factory,calendar)
 pending=service.generate_day(DAY,OPEN+timedelta(seconds=5))
 assert pending['status']=='pending' and 'not ended' in pending['error']
 unknown=daily.DailyReportService(factory,ExchangeCalendar({'rules':[]})).generate_day(DAY,AFTER)
 assert unknown['status']=='pending' and 'No verified' in unknown['error']
 with factory() as db:
  assert db.scalar(select(func.count()).select_from(daily.DailyReportPart))==0
  with pytest.raises(daily.DailyReportError):daily.artifact(db,DAY)


def test_long_held_trade_continues_next_date_with_stable_identifier(factory,calendar):
 with factory() as db:
  trade=insert_trade(db,created=OPEN-timedelta(days=25))
  for second in range(4):add_snapshot(db,trade,OPEN+timedelta(days=1,seconds=second))
  db.commit()
 service=make(factory,calendar);service.generate_day(DAY,AFTER)
 service.generate_day('2026-10-27',AFTER+timedelta(days=1))
 with factory() as db:
  first=daily.report_manifest(db,DAY)['manifest'];second=daily.report_manifest(db,'2026-10-27')['manifest']
  a,b=first['coverage'][0],second['coverage'][0]
  assert a['trade_id']==b['trade_id']=='t001' and a['strategy_id']==b['strategy_id']=='s001'
  assert a['sheet']==b['sheet'] and b['sheet'].endswith('_Continued')
  assert first['snapshot_count']==3 and second['snapshot_count']==4
  assert db.get(Trade,'t001').status=='open'


def test_ist_boundary_uses_utc_seconds_without_mixing_trading_dates(factory,calendar):
 with factory() as db:
  trade=insert_trade(db,seconds=0,created=OPEN-timedelta(days=1))
  for at in (datetime(2026,10,25,18,29,59,tzinfo=UTC),datetime(2026,10,25,18,30,tzinfo=UTC),OPEN):add_snapshot(db,trade,at)
  db.commit()
 service=make(factory,calendar);assert service.discover_dates()==2
 service.generate_day('2026-10-25',AFTER);service.generate_day(DAY,AFTER)
 with factory() as db:
  assert db.get(daily.DailyReport,'2026-10-25').snapshot_count==1
  assert db.get(daily.DailyReport,DAY).snapshot_count==2


def test_row_and_cell_limits_split_without_omitting_multileg_seconds(factory,calendar):
 with factory() as db:insert_trade(db,seconds=7,legs=3)
 service=make(factory,calendar,max_sheet_rows=11,max_data_rows_per_sheet=3,max_workbook_cells=1100)
 result=service.generate_day(DAY,AFTER)
 assert result['record_count']==21 and result['snapshot_count']==7
 with factory() as db:
  manifest=daily.report_manifest(db,DAY)['manifest'];parts=[a for a in manifest['artifacts'] if a['kind']=='xlsx']
  assert len(parts)>1 and any(a['kind']=='index_xlsx' for a in manifest['artifacts'])
  assert sum(p['record_count'] for p in parts)==21
  assert len({(c['part_number'],c['sheet']) for c in manifest['coverage']})==len(manifest['coverage'])
  all_rows=[]
  for part in parts:
   item=daily.artifact(db,DAY,'csv_gz',part['part_number'])
   all_rows.extend(csv.DictReader(io.StringIO(gzip.decompress(item['file_bytes']).decode('utf-8-sig'))))
  assert len(all_rows)==21
  assert len({(r['timestamp_utc'],r['leg_id']) for r in all_rows})==21


def test_byte_limit_triggers_automatic_smaller_parts(factory,calendar,monkeypatch):
 with factory() as db:insert_trade(db,seconds=12,legs=2)
 service=make(factory,calendar,max_workbook_cells=2500)
 service.limits=replace(service.limits,max_workbook_bytes=10000)
 original=daily._normalize_xlsx;calls=0
 def once(payload):
  nonlocal calls
  calls+=1
  content=original(payload)
  return content+b'padding'*2000 if calls==1 else content
 monkeypatch.setattr(daily,'_normalize_xlsx',once)
 result=service.generate_day(DAY,AFTER)
 assert result['status']=='finalized' and result['record_count']==24 and calls>1


def test_finalized_exports_are_immutable_and_generation_is_idempotent(factory,calendar):
 with factory() as db:insert_trade(db)
 service=make(factory,calendar);first=service.generate_day(DAY,AFTER)
 with factory() as db:
  original=daily.artifact(db,DAY)['sha256'];first_parts=db.scalar(select(func.count()).select_from(daily.DailyReportPart))
 second=service.generate_day(DAY,AFTER+timedelta(days=10))
 assert first['manifest_checksum']==second['manifest_checksum']
 with factory() as db:
  assert daily.artifact(db,DAY)['sha256']==original
  assert db.scalar(select(func.count()).select_from(daily.DailyReportPart))==first_parts
  assert db.get(daily.DailyReport,DAY).attempts==1


def test_only_explicit_matching_local_ack_is_accepted(factory,calendar):
 with factory() as db:insert_trade(db)
 make(factory,calendar).generate_day(DAY,AFTER)
 with factory() as db:
  report=db.get(daily.DailyReport,DAY)
  for confirmed,checksum,count in ((False,report.manifest_checksum,6),(True,'bad',6),(True,report.manifest_checksum,5)):
   with pytest.raises(daily.DailyReportError):daily.verify_local_ack(db,DAY,checksum,confirmed,count)
  daily.artifact(db,DAY)
  assert report.local_ack_at is None
  value=daily.verify_local_ack(db,DAY,report.manifest_checksum,True,6)
  assert value['local_ack_at'] and value['local_ack_checksum']==report.manifest_checksum
  assert db.get(Trade,'t001').status=='open'


def test_cloud_artifacts_verify_every_part_and_detect_corrupt_bytes(factory,calendar):
 with factory() as db:insert_trade(db)
 make(factory,calendar).generate_day(DAY,AFTER)
 with factory() as db:
  items=list(daily.cloud_artifacts(db,DAY))
  assert {i['kind'] for i in items}=={'xlsx','csv_gz'}
  for i in items:
   assert hashlib.sha256(i['file_bytes']).hexdigest()==i['sha256']
   assert hashlib.md5(i['file_bytes']).hexdigest()==i['md5']
  row=db.get(daily.DailyReportPart,f'{DAY}:xlsx:0001');row.file_bytes=b'corrupted';db.commit()
  with pytest.raises(daily.DailyReportError,match='integrity'):daily.artifact(db,DAY)


def test_no_quote_trade_gets_own_empty_sheet_and_gaps_are_explicit(factory,calendar):
 with factory() as db:
  insert_trade(db,seconds=0)
  db.add(RecordingGap(trade_id='t001',start_second=int(OPEN.timestamp()),end_second=int(OPEN.timestamp())+30,
   detected_at=AFTER.isoformat(),reason='recording_interrupted',data={'note':'No backfill'}));db.commit()
 result=make(factory,calendar).generate_day(DAY,AFTER)
 assert result['trade_count']==1 and result['record_count']==result['snapshot_count']==0
 with factory() as db:
  manifest=daily.report_manifest(db,DAY)['manifest']
  assert len(manifest['coverage'])==1 and len(manifest['recording_gaps'])==1
  wb=load_workbook(io.BytesIO(daily.artifact(db,DAY)['file_bytes']),read_only=True,data_only=True)
  assert len(wb.sheetnames)==2;wb.close()


def test_restart_discovers_unfinished_dates_and_reclaims_expired_lease(factory,calendar):
 with factory() as db:
  insert_trade(db);daily.ensure_day(db,DAY);db.commit()
  report=db.get(daily.DailyReport,DAY);report.status='generating';report.lease_owner='dead-process';report.lease_until=1;db.commit()
 service=make(factory,calendar)
 assert service.discover_dates()==0
 assert service.generate_day(DAY,AFTER)['status']=='finalized'
 with factory() as db:assert db.scalar(select(func.count()).select_from(PnlSnapshot))==3


def test_active_generation_lease_prevents_duplicate_export_worker(factory,calendar):
 import time
 with factory() as db:
  insert_trade(db);report=daily.ensure_day(db,DAY);report.status='generating';report.lease_owner='other';report.lease_until=time.time()+100;db.commit()
 result=make(factory,calendar).generate_day(DAY,AFTER)
 assert result['status']=='generating'
 with factory() as db:assert db.scalar(select(func.count()).select_from(daily.DailyReportPart))==0


def test_invalid_limits_and_naive_times_fail_before_export(factory,calendar):
 with pytest.raises(daily.DailyReportError):make(factory,calendar,max_sheet_rows=8)
 with pytest.raises(daily.DailyReportError):make(factory,calendar,max_workbook_cells=200)
 with pytest.raises(daily.DailyReportError):make(factory,calendar).generate_day(DAY,datetime(2026,10,26))


@pytest.mark.parametrize('day,at', [
 ('2026-10-24',datetime(2026,10,24,12,tzinfo=UTC)),  # Weekend: no exchange session.
 (DAY,AFTER),  # Weekday: the regular exchange session has already ended.
])
def test_demo_day_waits_for_ist_midnight_and_exports_later_seconds(factory,calendar,day,at):
 with factory() as db:
  trade=insert_trade(db,day=at,seconds=1)
  trade['demo']=True
  db.get(Trade,trade['id']).data=trade
  db.commit()
 service=make(factory,calendar)
 pending=service.generate_day(day,at+timedelta(seconds=2))
 assert pending['status']=='pending' and 'IST date ends' in pending['error']
 with factory() as db:
  assert db.scalar(select(func.count()).select_from(daily.DailyReportPart))==0
  add_snapshot(db,trade,at+timedelta(seconds=3));db.commit()
 midnight=datetime.fromtimestamp(daily.bounds(day)[1],UTC)
 assert service.generate_day(day,midnight-timedelta(seconds=1))['status']=='pending'
 finalized=service.generate_day(day,midnight)
 assert finalized['status']=='finalized' and finalized['snapshot_count']==2
 assert finalized['record_count']==4


def test_previous_demo_day_needs_no_real_calendar_but_live_trade_still_does(factory):
 unknown=ExchangeCalendar({'rules':[]})
 with factory() as db:
  trade=insert_trade(db,seconds=1)
  trade['demo']=True;db.get(Trade,trade['id']).data=trade;db.commit()
 service=make(factory,unknown)
 result=service.generate_day(DAY,AFTER+timedelta(days=1))
 assert result['status']=='finalized' and result['snapshot_count']==1
 # A real market-data trade on a different date must still fail closed.
 next_day=OPEN+timedelta(days=1)
 with factory() as db:insert_trade(db,identifier='live',day=next_day,seconds=1)
 result=service.generate_day('2026-10-27',AFTER+timedelta(days=2))
 assert result['status']=='pending' and 'No verified exchange calendar' in result['error']
