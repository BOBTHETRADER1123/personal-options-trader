"""Isolated recording/recovery tests; no user database or remote service access."""
import copy
import csv
import gzip
import hashlib
import io
from contextlib import closing
from pathlib import Path
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select, update, inspect, event as sqlalchemy_event
from sqlalchemy.orm import sessionmaker

from app import engine, reporting
from app.database import Base, configure_sqlite
from app.database import iso_now
from fastapi import HTTPException
from app.models import Event, PnlSnapshot, RecordingCursor, RecordingGap, SchedulerLease, Setting, Trade
from app.scheduler import CollectorLease, ExchangeCalendar, acquire_or_renew, database_seconds, guard_lease

UTC=timezone.utc
MOMENT=datetime(2026,10,23,3,46,0,tzinfo=UTC)  # Friday, 09:16 IST


def calendar(**changes):
    rule={'exchange':'NSE','segments':['nse_fo','D'],'valid_from':'2026-01-01',
          'valid_through':'2026-12-31','verified_at':'2026-10-23',
          'sources':['https://exchange.example/official-test-fixture'],
          'sessions':[['09:15','15:40']],'holidays':[]}
    rule.update(changes)
    return ExchangeCalendar({'rules':[rule]})


def trade(at=MOMENT,identifier='recording_test',demo=False):
    quote_at=(at-timedelta(seconds=2)).isoformat()
    return {'id':identifier,'broker':'kotak','strategy_id':'same_strategy','name':'Recorded call',
            'underlying':'NIFTY','status':'open','demo':demo,'created_at':at.isoformat(),
            'net_pnl':50.0,'gross_pnl':55.0,'unrealized_pnl':55.0,'realized_pnl':0.0,'charges':5.0,
            'legs':[{'id':'leg_a','instrument_id':'kotak:nse_fo:123','symbol':'TEST_CE','exchange':'NSE',
                     'segment':'nse_fo','underlying':'NIFTY','strike':23000,'expiry':'2026-10-27',
                     'option_type':'CE','side':'BUY','lots':1,'filled_lots':1,'open_lots':1,'lot_size':50,
                     'price_multiplier':1,'entry_price':100.0,'current_price':101.1,'realized_pnl':0.0,
                     'unrealized_pnl':55.0,'gross_pnl':55.0,'net_pnl':50.0,'charges':5.0,'status':'open',
                     'quote_at':quote_at,'quote_event_id':101,'quote_source':'kotak','bid':101,'ask':101.2}]}


@pytest.fixture
def journal(tmp_path):
    sql=create_engine(f'sqlite:///{tmp_path / "recording.db"}')
    configure_sqlite(sql)
    Base.metadata.create_all(sql)
    factory=sessionmaker(sql,expire_on_commit=False)
    engine.clear_cache()
    with factory() as db:
        engine.update_settings(db,{'logging_mode':'full_history','demo_mode':False})
        db.commit()
    yield factory
    engine.clear_cache()
    sql.dispose()


def put(db,value):
    engine.store_trade(db,value)
    db.commit()


def test_second_key_is_idempotent_and_reused_quotes_keep_original_time(journal):
    with journal() as db:
        value=trade();put(db,value)
        assert engine.take_snapshots(db,now=MOMENT,calendar=calendar())['recorded']==1
        db.commit()
        assert engine.take_snapshots(db,now=MOMENT+timedelta(microseconds=900000),calendar=calendar())['recorded']==0
        assert engine.take_snapshots(db,now=MOMENT+timedelta(seconds=1),calendar=calendar())['recorded']==1
        db.commit()
        rows=db.scalars(select(PnlSnapshot).order_by(PnlSnapshot.snapshot_second)).all()
        assert len(rows)==2
        first,last=rows[0].data,rows[1].data
        assert first['legs'][0]['quote_at']==last['legs'][0]['quote_at']==value['legs'][0]['quote_at']
        assert last['legs'][0]['quote_reused'] is True and last['carried_forward'] is True
        assert last['legs'][0]['quote_age_seconds']==3
        assert last['timestamp_ist'].endswith('+05:30')
        assert last['legs'][0]['lot_size']==50 and last['legs'][0]['strategy_id']=='same_strategy'


def test_outage_is_a_gap_not_fabricated_rows_and_marks_stale(journal):
    with journal() as db:
        put(db,trade())
        engine.take_snapshots(db,now=MOMENT,calendar=calendar());db.commit()
        engine.clear_cache(db)  # Separate process/restart restores permanent trade/cursor.
        engine.take_snapshots(db,now=MOMENT+timedelta(seconds=20),calendar=calendar());db.commit()
        assert db.scalar(select(func.count()).select_from(PnlSnapshot))==2
        gap=db.scalar(select(RecordingGap))
        assert gap.end_second-gap.start_second+1==19
        last=db.scalar(select(PnlSnapshot).order_by(PnlSnapshot.snapshot_second.desc()))
        assert last.data['legs'][0]['data_freshness']=='stale'
        assert last.data['valuation_stale'] is True
        assert 'reconstructed' in gap.data['note']


def test_no_zero_position_records_and_unknown_calendar_fails_closed(journal):
    with journal() as db:
        assert engine.take_snapshots(db,now=MOMENT,calendar=calendar())['recorded']==0
        value=trade();value['status']='pending';value['legs'][0]['open_lots']=0;put(db,value)
        assert engine.take_snapshots(db,now=MOMENT,calendar=calendar())['recorded']==0
        value['status']='open';value['legs'][0]['open_lots']=1;put(db,value)
        result=engine.take_snapshots(db,now=MOMENT,calendar=ExchangeCalendar({'rules':[]}))
        assert result['calendar_unknown']==1 and result['recorded']==0
        assert db.scalar(select(func.count()).select_from(PnlSnapshot))==0


def test_gaps_exclude_overnight_weekends_and_holidays(journal):
    friday=datetime(2026,10,23,3,45,58,tzinfo=UTC)
    monday=datetime(2026,10,26,3,45,2,tzinfo=UTC)
    maintained=calendar(sessions=[['09:15','09:16']])
    with journal() as db:
        put(db,trade(friday))
        engine.take_snapshots(db,now=friday,calendar=maintained);db.commit()
        engine.take_snapshots(db,now=monday,calendar=maintained);db.commit()
        gaps=db.scalars(select(RecordingGap).order_by(RecordingGap.start_second)).all()
        assert len(gaps)==2 and sum(g.end_second-g.start_second+1 for g in gaps)==3
        assert maintained.state(trade()['legs'][0],datetime(2026,10,24,4,tzinfo=UTC))['open'] is False
        assert calendar(holidays=['2026-10-23']).state(trade()['legs'][0],MOMENT)['open'] is False


def test_special_session_overrides_weekend_and_commodity_rules_are_specific():
    sunday=datetime(2026,11,8,12,30,tzinfo=UTC)
    maintained=calendar(holidays=['2026-11-08'],overrides={'2026-11-08':[['18:00','19:00']]})
    assert maintained.state({**trade()['legs'][0],'expiry':'2026-11-24'},sunday)['open'] is True
    mcx=copy.deepcopy(maintained.rules[0]);mcx.update(exchange='MCX',segments=['mcx_fo'],underlyings=['GOLD'],sessions=[['09:00','23:55']])
    maintained=ExchangeCalendar({'rules':[mcx]})
    assert maintained.state({'exchange':'MCX','segment':'mcx_fo','underlying':'GOLD'},MOMENT)['open'] is True
    assert maintained.state({'exchange':'MCX','segment':'mcx_fo','underlying':'COTTON'},MOMENT)['known'] is False


def test_database_lease_excludes_second_collector_and_fences_old_generation(journal):
    with journal() as db:
        first=acquire_or_renew(db,'process_a');db.commit()
        assert first is not None
        assert acquire_or_renew(db,'process_b') is None;db.commit()
        assert guard_lease(db,first) is True;db.commit()
        db.execute(update(SchedulerLease).values(expires_at=database_seconds(db)-1));db.commit()
        second=acquire_or_renew(db,'process_b');db.commit()
        assert second.generation==first.generation+1
        assert guard_lease(db,first) is False
        assert guard_lease(db,second) is True
        put(db,trade())
        assert engine.take_snapshots(db,now=MOMENT,calendar=calendar(),lease=first)['lease_lost'] is True
        assert engine.take_snapshots(db,now=MOMENT,calendar=calendar(),lease=second)['recorded']==1
        db.commit()


def test_collector_release_allows_replacement_without_deleting_rows(journal):
    a,b=CollectorLease(journal),CollectorLease(journal)
    assert a.renew() is True and b.renew() is False
    a.release()
    assert b.renew() is True
    b.release()
    with journal() as db:
        assert db.scalar(select(func.count()).select_from(SchedulerLease))==1


def test_quote_buffer_coalesces_and_dequeues_only_after_committed_capture(journal):
    with journal() as db:
        value=trade();put(db,value)
        for price in (101,102,103):
            engine.buffer_quote({'broker':'kotak','instrument_id':value['legs'][0]['instrument_id'],
                                 'ltp':price,'received_at':MOMENT.isoformat(),'demo':False})
        assert engine.flush_buffered_quotes(db)==1
        db.rollback()
        assert engine.buffered_quote_status()['queued_contracts']==1
        assert engine.flush_buffered_quotes(db)==1
        db.commit()
        assert engine.buffered_quote_status()['queued_contracts']==0
        assert db.get(Trade,value['id']).data['legs'][0]['current_price']==103


def test_compact_excel_and_separate_full_csv_with_checksums(journal):
    with journal() as db:
        value=trade();put(db,value)
        engine.take_snapshots(db,now=MOMENT,calendar=calendar());db.commit()
        engine.take_snapshots(db,now=MOMENT+timedelta(seconds=4),calendar=calendar())
        db.add(Event(trade_id=value['id'],timestamp=MOMENT.isoformat(),kind='entry',data={'broker':'kotak'}))
        value.update(status='closed',closed_at=(MOMENT+timedelta(seconds=5)).isoformat(),history_available=True)
        job=reporting.enqueue_export(db,value,[],[],[]);identity=job.id;db.commit()
    assert reporting.process_pending_once(journal)==1
    with journal() as db:
        artifacts=reporting.get_export_artifacts(db,identity)
        assert {item['kind'] for item in artifacts}=={'xlsx','csv_gz'}
        for artifact in artifacts:
            assert artifact['size']==len(artifact['file_bytes'])
            assert artifact['sha256']==hashlib.sha256(artifact['file_bytes']).hexdigest()
        xlsx=next(a for a in artifacts if a['kind']=='xlsx')
        with closing(load_workbook(io.BytesIO(xlsx['file_bytes']))) as workbook:
            assert set(workbook.sheetnames)==reporting.SUMMARY_SHEETS
        archive=next(a for a in artifacts if a['kind']=='csv_gz')
        rows=list(csv.DictReader(io.StringIO(gzip.decompress(archive['file_bytes']).decode())))
        assert sum(row['record_type']=='one_second' for row in rows)==2
        assert {'trade_event','recording_gap'}.issubset({row['record_type'] for row in rows})
        snapshot=next(row for row in rows if row['record_type']=='one_second')
        assert snapshot['instrument_id']==value['legs'][0]['instrument_id']
        assert snapshot['quote_at']==value['legs'][0]['quote_at']
        assert snapshot['lot_size']=='50'


def test_local_excel_ack_does_not_delete_unverified_full_history(journal):
    with journal() as db:
        value=trade();put(db,value)
        engine.take_snapshots(db,now=MOMENT,calendar=calendar())
        value.update(status='closed',closed_at=MOMENT.isoformat(),history_available=True)
        job=reporting.enqueue_export(db,value,[],[],[]);db.commit()
        job.status='saved';job.verification_status='verified';job.checksum='x'*64;job.saved_at=MOMENT.isoformat()
        reporting.configure_archival(lambda db:True,lambda db,trade_id:(_ for _ in ()).throw(AssertionError('must not delete')))
        try:
            reporting.archive_after_verified_save(db,job);db.commit()
            assert job.deletion_status=='retained_until_drive_verified'
            assert db.scalar(select(func.count()).select_from(PnlSnapshot))==1
        finally:reporting.configure_archival(lambda db:False)


def test_additive_migration_preserves_existing_data(journal):
    from app.maintenance import migrate
    with journal() as db:
        put(db,trade())
        target=db.get_bind()
    migrate(target);migrate(target)
    with journal() as db:
        assert db.get(Trade,'recording_test') is not None
        from app.models import SchemaVersion
        assert db.get(SchemaVersion,2) is not None


def test_old_recording_policy_upgrades_once_without_removing_history(journal):
    with journal() as db:
        put(db,trade())
        setting=db.get(Setting,'app');setting.value={'logging_mode':'summary','recording_policy_version':2}
        db.commit();engine.clear_cache(db)
        assert engine.get_settings(db)['logging_mode']=='full_history'
        engine.update_settings(db,{'logging_mode':'summary'});db.commit();engine.clear_cache(db)
        assert engine.get_settings(db)['logging_mode']=='summary'
        assert db.get(Trade,'recording_test') is not None


def test_full_history_dashboard_uses_bounded_chart_cache_not_raw_history_reads(journal):
    with journal() as db:
        put(db,trade())
        engine.take_snapshots(db,now=MOMENT,calendar=calendar());db.commit()
        first=engine.dashboard(db)
        assert first['recording_mode']=='full_history'
        assert first['portfolio_timeline'][-1]['net_pnl']==50
        statements=[]
        def seen(connection,cursor,statement,params,context,many):statements.append(statement)
        sqlalchemy_event.listen(db.get_bind(),'before_cursor_execute',seen)
        try:
            for _ in range(5):assert engine.dashboard(db)['portfolio_timeline'][-1]['net_pnl']==50
            assert statements==[]
        finally:sqlalchemy_event.remove(db.get_bind(),'before_cursor_execute',seen)


def test_buffer_does_not_drop_a_new_quote_arriving_during_commit(journal):
    with journal() as db:
        value=trade();put(db,value)
        quote={'broker':'kotak','instrument_id':value['legs'][0]['instrument_id'],
               'ltp':101,'received_at':MOMENT.isoformat(),'demo':False}
        engine.buffer_quote(quote);engine.flush_buffered_quotes(db)
        engine.buffer_quote({**quote,'ltp':105});db.commit()
        assert engine.buffered_quote_status()['queued_contracts']==1
        engine.flush_buffered_quotes(db);db.commit()
        assert engine.buffered_quote_status()['queued_contracts']==0
        assert db.get(Trade,value['id']).data['legs'][0]['current_price']==105


def test_empty_report_queue_permits_neon_to_sleep(journal):
    assert reporting.next_work_delay(journal)>=600


def create_test_strategy(db,**settings):
    """Documented deterministic contract fixtures, explicitly not broker data."""
    contracts=[{'id':f'kotak:nse_fo:{index}','instrument_id':f'kotak:nse_fo:{index}',
                'broker':'kotak','exchange':'NSE','segment':'nse_fo','underlying':'NIFTY',
                'symbol':f'TEST_{index}_CE','strike':23000+index*50,'expiry':'2027-01-01',
                'option_type':'CE','lot_size':50,'price_multiplier':1,'spot':23000,'demo':False}
               for index in (1,2)]
    class FixtureCatalog:
        def find(self,identity,*args):return next((c for c in contracts if c['id']==identity),None)
        def list_contracts(self,**kwargs):return contracts
    engine.set_catalog(FixtureCatalog())
    engine.update_settings(db,{'demo_mode':False,'brokerage_per_order':0,'charges_percent':0,'fill_basis':'ltp',**settings})
    for contract in contracts:
        engine.record_quote(db,{'broker':'kotak','instrument_id':contract['id'],'ltp':100,
                                'received_at':iso_now(),'demo':False})
    result=engine.create_trades(db,{'name':'Fixture paper call','underlying':'NIFTY','brokers':['kotak'],
        'legs':[{'instrument_id':contracts[0]['id'],'side':'BUY','lots':1,'enabled':True}],
        'idempotency_key':'fixture-open-strategy','confirmed':True})
    db.commit()
    return result,contracts


def test_replacement_closes_and_reopens_with_original_entry_preserved(journal):
    with journal() as db:
        result,contracts=create_test_strategy(db)
        value=result['trades'][0];old=value['legs'][0]
        engine.record_quote(db,{'broker':'kotak','instrument_id':contracts[0]['id'],'ltp':120,'received_at':iso_now(),'demo':False})
        engine.record_quote(db,{'broker':'kotak','instrument_id':contracts[1]['id'],'ltp':80,'received_at':iso_now(),'demo':False})
        db.commit()
        result=engine.apply_action(db,value['id'],{'action':'replace_leg','leg_id':old['id'],
            'leg':{'instrument_id':contracts[1]['id'],'side':'SELL','lots':1},
            'idempotency_key':'fixture-replace-leg','confirmed':True});db.commit()
        value=result['trade']
        assert value['status']=='open' and len(value['legs'])==2
        assert value['legs'][0]['entry_price']==100 and value['legs'][0]['exit_price']==120
        assert value['legs'][0]['open_lots']==0 and value['legs'][1]['entry_price']==80
        assert value['realized_pnl']==1000
        kinds=set(db.scalars(select(Event.kind).where(Event.trade_id==value['id'])))
        assert {'replacement_close','replacement_open'}.issubset(kinds)


def test_real_data_entry_assumption_edit_is_rejected(journal):
    with journal() as db:
        result,contracts=create_test_strategy(db)
        value=result['trades'][0]
        with pytest.raises(HTTPException,match='cannot be edited'):
            engine.apply_action(db,value['id'],{'action':'edit','leg_id':value['legs'][0]['id'],'price':1,
                'idempotency_key':'fixture-edit-price','confirmed':True})
        db.rollback()
        assert db.get(Trade,value['id']).data['legs'][0]['entry_price']==100


def test_finite_simulated_capital_and_unlimited_option_keep_margin(journal):
    with journal() as db:
        result,contracts=create_test_strategy(db,paper_capital=1)
        assert result['trades']==[] and 'simulated capital' in result['results'][0]['error']
        engine.update_settings(db,{'unlimited_paper_buying_power':True});db.commit()
        engine.record_quote(db,{'broker':'kotak','instrument_id':contracts[0]['id'],'ltp':100,'received_at':iso_now(),'demo':False})
        result=engine.create_trades(db,{'name':'Unlimited paper only','underlying':'NIFTY','brokers':['kotak'],
            'legs':[{'instrument_id':contracts[0]['id'],'side':'BUY','lots':1,'enabled':True}],
            'idempotency_key':'fixture-unlimited-open','confirmed':True});db.commit()
        assert result['trades'][0]['margin_estimate']==5000


def test_verified_pruning_gate_retains_all_source_rows_until_drive_matches(journal):
    with journal() as db:
        value=trade();put(db,value)
        engine.take_snapshots(db,now=MOMENT,calendar=calendar())
        value.update(status='closed',closed_at=MOMENT.isoformat(),history_available=True)
        engine.store_trade(db,value)
        job=reporting.enqueue_export(db,value,[],[],[]);identity=job.id;db.commit()
    assert reporting.process_pending_once(journal)==1
    with journal() as db:
        with pytest.raises(ValueError,match='verified Drive manifest'):
            reporting.prune_verified_snapshots(db,identity)
        db.rollback()
        assert db.scalar(select(func.count()).select_from(PnlSnapshot))==1


def test_alembic_upgrades_an_existing_journal_and_is_idempotent(tmp_path):
    from alembic import command
    from alembic.config import Config
    sql=create_engine(f'sqlite:///{tmp_path / "old-journal.db"}')
    configure_sqlite(sql)
    Trade.__table__.create(sql)
    with sql.begin() as connection:
        connection.execute(Trade.__table__.insert(),{'id':'legacy_before_migration','broker':'dhan',
            'status':'closed','data':{'must_survive':'all historical records'}})
    configuration=Config(str(Path(__file__).resolve().parents[1] / 'alembic.ini'))
    try:
        with sql.begin() as connection:
            configuration.attributes['connection']=connection
            command.upgrade(configuration,'head')
            command.upgrade(configuration,'head')
        with sql.connect() as connection:
            assert connection.scalar(select(Trade.data))=={'must_survive':'all historical records'}
            assert connection.scalar(select(func.count()).select_from(PnlSnapshot))==0
            assert {'pnl_snapshots','scheduler_leases','cloud_backup_jobs','export_artifacts'}.issubset(inspect(connection).get_table_names())
    finally:sql.dispose()


def test_additive_alembic_transaction_rolls_back_without_touching_old_rows(tmp_path):
    from alembic import command
    from alembic.config import Config
    sql=create_engine(f'sqlite:///{tmp_path / "rollback-journal.db"}')
    configure_sqlite(sql);Trade.__table__.create(sql)
    with sql.begin() as connection:
        connection.execute(Trade.__table__.insert(),{'id':'preserved','broker':'kotak','status':'closed','data':{'net_pnl':123}})
    configuration=Config(str(Path(__file__).resolve().parents[1] / 'alembic.ini'))
    try:
        with pytest.raises(RuntimeError,match='simulated interruption'):
            with sql.begin() as connection:
                configuration.attributes['connection']=connection
                command.upgrade(configuration,'head')
                raise RuntimeError('simulated interruption')
        with sql.connect() as connection:
            assert connection.scalar(select(Trade.data))=={'net_pnl':123}
            assert 'pnl_snapshots' not in inspect(connection).get_table_names()
    finally:sql.dispose()


@pytest.mark.parametrize('demo_enabled,include_real,market_open,expected', [
    (False,False,False,False),
    (True,False,False,True),
    (False,True,True,True),
    (True,True,True,True),
    (False,True,False,False),
])
def test_recording_demand_ignores_disabled_samples_and_keeps_real_sessions(monkeypatch,demo_enabled,include_real,market_open,expected):
    from app import main
    sample=trade(identifier='sample-demand',demo=True)
    real=trade(identifier='real-demand')
    calls=[]
    class DemandCalendar:
        def state(self,leg,now):
            calls.append(leg)
            return {'open':market_open,'known':True}
    monkeypatch.setattr(main,'_runtime_settings',{'demo_mode':demo_enabled})
    monkeypatch.setattr(main,'_active_positions',([real] if include_real else [])+[sample])
    monkeypatch.setattr(main,'exchange_calendar',DemandCalendar())
    assert main.recording_needed() is expected
    # Synthetic contracts must never keep a genuine-mode collector alive by
    # falling through to the real exchange calendar.
    assert len(calls)==(1 if include_real else 0)


def test_sample_off_stops_new_snapshots_without_erasing_cursor_and_resumes_with_gap(journal):
    with journal() as db:
        sample=trade(identifier='sample-pause',demo=True)
        engine.update_settings(db,{'demo_mode':True});put(db,sample)
        assert engine.take_snapshots(db,now=MOMENT,calendar=calendar())['recorded']==1
        db.commit()
        engine.update_settings(db,{'demo_mode':False});db.commit()
        original=db.get(Trade,sample['id']).data.copy()
        statements=[]
        def writes(connection,cursor,statement,params,context,many):
            if statement.lstrip().split(None,1)[0].upper() in {'INSERT','UPDATE','DELETE','REPLACE'}:statements.append(statement)
        sqlalchemy_event.listen(db.get_bind(),'before_cursor_execute',writes)
        try:
            assert engine.take_snapshots(db,now=MOMENT+timedelta(seconds=5),calendar=calendar())['recorded']==0
            db.commit()
            assert statements==[]
        finally:sqlalchemy_event.remove(db.get_bind(),'before_cursor_execute',writes)
        assert db.scalar(select(func.count()).select_from(PnlSnapshot))==1
        assert db.get(RecordingCursor,sample['id']).last_second==int(MOMENT.timestamp())
        assert db.get(Trade,sample['id']).data==original
        engine.update_settings(db,{'demo_mode':True});db.commit()
        assert engine.take_snapshots(db,now=MOMENT+timedelta(seconds=10),calendar=calendar())['recorded']==1
        db.commit()
        assert db.scalar(select(func.count()).select_from(PnlSnapshot))==2
        gap=db.scalar(select(RecordingGap).where(RecordingGap.trade_id==sample['id']))
        assert gap.start_second==int(MOMENT.timestamp())+1 and gap.end_second==int(MOMENT.timestamp())+9
        assert gap.reason=='recording_interrupted' and 'reconstructed' in gap.data['note']


@pytest.mark.parametrize('mode',['full_history','every_update'])
@pytest.mark.parametrize('demo_enabled',[False,True])
def test_real_positions_record_independently_of_sample_view_in_both_recording_modes(journal,mode,demo_enabled):
    from app.models import Timeline
    with journal() as db:
        engine.update_settings(db,{'logging_mode':mode,'demo_mode':demo_enabled});db.commit()
        put(db,trade(identifier='sample',demo=True))
        put(db,trade(identifier='genuine',demo=False))
        result=engine.take_snapshots(db,now=MOMENT,calendar=calendar());db.commit()
        assert result['recorded']==(2 if demo_enabled else 1)
        model=PnlSnapshot if mode=='full_history' else Timeline
        identifiers=set(db.scalars(select(model.trade_id)))
        assert identifiers==({'genuine','sample'} if demo_enabled else {'genuine'})
        assert db.get(Trade,'sample') is not None


def test_legacy_sample_only_disabled_creates_no_writes(journal):
    from app.models import Timeline
    with journal() as db:
        engine.update_settings(db,{'logging_mode':'every_update','demo_mode':False});db.commit()
        put(db,trade(identifier='disabled-legacy',demo=True))
        statements=[]
        def writes(connection,cursor,statement,params,context,many):
            if statement.lstrip().split(None,1)[0].upper() in {'INSERT','UPDATE','DELETE','REPLACE'}:statements.append(statement)
        sqlalchemy_event.listen(db.get_bind(),'before_cursor_execute',writes)
        try:
            assert engine.take_snapshots(db,now=MOMENT,calendar=calendar())['recorded']==0
            db.commit();assert statements==[]
        finally:sqlalchemy_event.remove(db.get_bind(),'before_cursor_execute',writes)
        assert db.scalar(select(func.count()).select_from(Timeline))==0
        assert db.scalar(select(func.count()).select_from(PnlSnapshot))==0
        assert db.scalar(select(func.count()).select_from(RecordingCursor))==0
        assert db.get(Trade,'disabled-legacy') is not None
