from __future__ import annotations

import hashlib
import io
import json
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app import companion_api, reporting
from app.database import Base
from app.reporting import DailyExportReconciliation, ExportJob, build_workbook, enqueue_export

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from companion.companion import Companion, validate_backend_url


@pytest.fixture
def factory(tmp_path):
    engine = create_engine(f"sqlite:///{(tmp_path / 'reports.sqlite3').as_posix()}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    yield factory
    engine.dispose()


@pytest.fixture
def payload():
    return {
        "format": "detailed",  # Explicit legacy reports remain readable.
        "trade": {
            "id": "trade_001", "strategy_id": "strategy_001", "name": "Bull Call Spread", "broker": "Dhan",
            "underlying": "NIFTY", "status": "closed", "created_at": "2026-10-22T03:45:00+00:00",
            "closed_at": "2026-10-22T08:13:00+00:00", "gross_pnl": 1420, "net_pnl": 1400,
            "charges": 20, "realized_pnl": 1420, "unrealized_pnl": 0,
            "mfe": 2000, "mae": -1900, "mfe_at": "2026-10-22T07:13:00+00:00", "mae_at": "2026-10-22T05:01:00+00:00",
            "legs": [{"id": "leg_1", "symbol": "NIFTY_TEST", "instrument_id": "test:1", "side": "BUY", "lots": 1,
                      "lot_size": 10, "open_lots": 0, "entry_price": 10, "exit_price": 152, "realized_pnl": 1420,
                      "exchange": "NSE", "expiry": "2026-10-27", "strike": 23400, "option_type": "CE"}],
        },
        "events": [{"id": 1, "kind": "entry", "timestamp": "2026-10-22T03:45:00+00:00", "price": 10}],
        "timeline": [{"id": 1, "timestamp": "2026-10-22T07:13:00+00:00", "net_pnl": 2000, "carried_forward": False},
                     {"id": 2, "timestamp": "2026-10-22T07:13:01+00:00", "net_pnl": 2000, "carried_forward": True}],
        "quotes": [{"id": 1, "broker": "Dhan", "instrument_id": "test:1", "received_at": "2026-10-22T07:13:00.123456+00:00",
                    "received_ns": 1792653180123456789, "ltp": 210}],
    }


def enqueue(db, payload):
    job = enqueue_export(db, payload["trade"], payload["events"], payload["timeline"], payload["quotes"], report_format=payload.get("format", "summary"))
    db.commit()
    return job


def ready(factory, payload):
    with factory() as db:
        job_id = enqueue(db, payload).id
    assert reporting.process_pending_once(factory) == 1
    with factory() as db:
        return db.get(ExportJob, job_id)


def test_workbook_has_complete_history_and_ist_times(payload):
    data = build_workbook(payload)
    book = load_workbook(io.BytesIO(data), data_only=False)
    assert reporting.REQUIRED_SHEETS.issubset(book.sheetnames)
    values = dict(book["Trade Summary"].iter_rows(min_row=2, max_col=2, values_only=True))
    assert values["Closing time (IST)"].hour == 13
    assert values["Closing time (IST)"].minute == 43
    assert values["MFE time (IST)"].hour == 12
    assert values["MAE time (IST)"].hour == 10
    assert values["Profit giveback (INR)"] == 600
    assert book["P&L Timeline"].max_row == 3
    assert book["P&L Timeline"].cell(3, 14).value is True
    assert book["Market Updates"].cell(2, 2).value == "1792653180123456789"
    assert "received_ns" in book["Market Updates"].cell(2, 19).value
    book.close()


def test_formula_injection_is_escaped(payload):
    payload["trade"]["name"] = '=HYPERLINK("https://attacker.invalid","click")'
    book = load_workbook(io.BytesIO(build_workbook(payload)))
    row = next(row for row in book["Trade Summary"] if row[0].value == "Strategy name")
    assert row[1].data_type == "s"
    assert row[1].value.startswith("'=")
    book.close()


def test_sample_origin_is_explicit_in_report_and_job(factory, payload):
    payload["trade"]["demo"] = True
    with factory() as db:
        job = enqueue(db, payload)
        assert reporting.public_job(job)["demo"] is True
    book = load_workbook(io.BytesIO(build_workbook(payload)))
    summary = dict(book["Trade Summary"].iter_rows(min_row=2, max_col=2, values_only=True))
    assert "SYNTHETIC SAMPLE" in summary["Quote and contract origin"]
    book.close()


def test_mcx_report_preserves_multiplier_and_monetary_quantity(payload):
    payload["trade"]["legs"][0].update(exchange="MCX", lots=2, lot_size=1, price_multiplier=100, multiplier_verified=True)
    book = load_workbook(io.BytesIO(build_workbook(payload)))
    sheet = book["Strategy Legs"]
    headers = {cell.value:cell.column for cell in sheet[1]}
    assert sheet.cell(2,headers["Total contract quantity"]).value == 2
    assert sheet.cell(2,headers["Monetary price multiplier"]).value == 100
    assert sheet.cell(2,headers["Quantity used for P&L"]).value == 200
    assert sheet.cell(2,headers["Multiplier verified from master"]).value is True
    book.close()


def test_trade_and_broker_paths_are_unique(factory, payload):
    with factory() as db:
        first = enqueue(db, payload)
        second_payload = json.loads(json.dumps(payload))
        second_payload["trade"].update(id="trade_002", broker="KotakNeo")
        second = enqueue(db, second_payload)
        assert first.filename != second.filename
        assert first.relative_path.startswith("Dhan/2026/10/Trade_trade_001/")
        assert second.relative_path.startswith("KotakNeo/2026/10/Trade_trade_002/")


def test_enqueue_and_generation_are_idempotent(factory, payload):
    with factory() as db:
        first = enqueue(db, payload)
        second = enqueue(db, payload)
        assert first.id == second.id
        assert len(db.scalars(select(ExportJob)).all()) == 1
    assert reporting.process_pending_once(factory) == 1
    assert reporting.process_pending_once(factory) == 0


def test_active_trade_cannot_be_archived(factory, payload):
    payload["trade"]["status"] = "open"
    with factory() as db, pytest.raises(ValueError, match="completed"):
        enqueue(db, payload)


def test_restart_recovers_expired_generation_lease(factory, payload):
    with factory() as db:
        job = enqueue(db, payload)
        job.status, job.lease_until = "generating", "2000-01-01T00:00:00+00:00"
        db.commit()
    assert reporting.process_pending_once(factory) == 1
    with factory() as db:
        recovered = db.get(ExportJob, job.id)
        assert recovered.status == "ready" and recovered.file_bytes and recovered.payload["quotes"]


def test_generation_failure_retains_source_and_retry(factory, payload, monkeypatch):
    original = reporting.build_workbook
    with factory() as db:
        job = enqueue(db, payload)
    monkeypatch.setattr(reporting, "build_workbook", lambda _: (_ for _ in ()).throw(OSError("locked")))
    assert reporting.process_pending_once(factory) == 0
    with factory() as db:
        failed = db.get(ExportJob, job.id)
        assert failed.status == "failed"
        assert failed.payload["events"] == payload["events"]
        failed.lease_until = "2000-01-01T00:00:00+00:00"
        db.commit()
    monkeypatch.setattr(reporting, "build_workbook", original)
    assert reporting.process_pending_once(factory) == 1


def test_laptop_offline_keeps_ready_job_and_daily_reconciliation(factory, payload):
    job = ready(factory, payload)
    reporting.save_daily_reconciliation(factory)
    with factory() as db:
        assert db.get(ExportJob, job.id).status == "ready"
        daily = db.scalars(select(DailyExportReconciliation)).one()
        assert daily.data["pending"] == 1 and daily.data["saved"] == 0
        assert daily.data["mismatches"] == []


@pytest.fixture
def api(factory, monkeypatch):
    token = "companion-test-token-" + "a" * 40
    monkeypatch.setenv("COMPANION_TOKEN_SHA256", hashlib.sha256(token.encode()).hexdigest())
    app = FastAPI()
    app.include_router(companion_api.router_companion)
    app.include_router(companion_api.router_browser)
    def override():
        with factory() as db:
            yield db
    app.dependency_overrides[companion_api.db_session] = override
    reporting.configure_archival(lambda db: False)
    client = TestClient(app, base_url="https://testserver")
    client.headers["Authorization"] = "Bearer " + token
    yield client
    reporting.configure_archival(lambda db: False)
    client.close()


def claim_and_ack(api, job):
    claim = api.post(f"/api/companion/jobs/{job.id}/claim", json={"companion_id": "laptop_test_001"})
    assert claim.status_code == 200
    headers = {"X-Export-Claim": claim.json()["claim_token"]}
    ack = {"companion_id": "laptop_test_001", "checksum": job.checksum, "size": job.size,
           "filename": job.filename, "local_path": "C:/Trades/" + job.filename, "workbook_valid": True}
    return headers, ack


def test_companion_requires_authentication_and_tls(api):
    assert api.get("/api/companion/jobs", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert api.get("http://public.example/api/companion/jobs").status_code == 400


def test_wrong_ack_does_not_save_or_delete(api, factory, payload):
    job = ready(factory, payload)
    headers, ack = claim_and_ack(api, job)
    ack["checksum"] = "0" * 64
    response = api.post(f"/api/companion/jobs/{job.id}/ack", headers=headers, json=ack)
    assert response.status_code == 422
    with factory() as db:
        retained = db.get(ExportJob, job.id)
        assert retained.status == "claimed" and retained.payload["quotes"]
        assert retained.deletion_status == "retained"


def test_verified_ack_is_idempotent_and_deletion_defaults_off(api, factory, payload):
    job = ready(factory, payload)
    headers, ack = claim_and_ack(api, job)
    response = api.post(f"/api/companion/jobs/{job.id}/ack", headers=headers, json=ack)
    assert response.status_code == 200 and response.json()["status"] == "saved"
    repeated = api.post(f"/api/companion/jobs/{job.id}/ack", headers=headers, json=ack)
    assert repeated.status_code == 200 and repeated.json()["already_saved"]
    with factory() as db:
        retained = db.get(ExportJob, job.id)
        assert retained.deletion_status == "retained" and retained.payload["quotes"]


def test_optional_deletion_runs_only_after_verified_ack(api, factory, payload):
    called = []
    reporting.configure_archival(lambda db: True, lambda db, trade_id: called.append(trade_id) or {"events": 1, "timeline": 2})
    job = ready(factory, payload)
    headers, ack = claim_and_ack(api, job)
    assert called == []
    response = api.post(f"/api/companion/jobs/{job.id}/ack", headers=headers, json=ack)
    assert response.status_code == 200
    assert called == ["trade_001"]
    assert response.json()["deletion_status"] == "archived"
    with factory() as db:
        archived = db.get(ExportJob, job.id)
        assert archived.file_bytes and archived.checksum and archived.payload["archived"]


def test_browser_download_never_marks_local_save(api, factory, payload):
    job = ready(factory, payload)
    response = api.get(f"/api/exports/{job.id}/download")
    assert response.status_code == 200
    assert hashlib.sha256(response.content).hexdigest() == job.checksum
    with factory() as db:
        assert db.get(ExportJob, job.id).status == "ready"


def test_local_save_failure_is_visible_and_retries_same_workbook(api, factory, payload):
    job = ready(factory, payload)
    headers, _ = claim_and_ack(api, job)
    response = api.post(f"/api/companion/jobs/{job.id}/failed", headers=headers, json={"message": "Workbook is locked"})
    assert response.status_code == 200 and response.json()["status"] == "failed"
    with factory() as db:
        retained = db.get(ExportJob, job.id)
        assert retained.payload["quotes"] and retained.file_bytes and retained.saved_at is None
        retained.lease_until = "2000-01-01T00:00:00+00:00"
        db.commit()
    assert reporting.process_pending_once(factory) == 0
    candidates = api.get("/api/companion/jobs").json()["jobs"]
    assert candidates[0]["id"] == job.id
    reclaimed = api.post(f"/api/companion/jobs/{job.id}/claim", json={"companion_id": "laptop_test_001"})
    assert reclaimed.status_code == 200


def companion_job(payload):
    content = build_workbook(payload)
    return {"id": "job_1", "filename": "NIFTY_trade_001.xlsx", "relative_path": "Dhan/2026/10/Trade_trade_001/NIFTY_trade_001.xlsx",
            "checksum": hashlib.sha256(content).hexdigest(), "size": len(content), "claim_token": "claim-token"}, content


def test_companion_restarts_after_save_before_ack_without_redownload(tmp_path, payload):
    job, content = companion_job(payload)
    counts = {"download": 0, "ack": 0}
    def handler(request):
        if request.url.path.endswith("/download"):
            counts["download"] += 1
            return httpx.Response(200, content=content)
        if request.url.path.endswith("/ack"):
            counts["ack"] += 1
            if counts["ack"] == 1:
                raise httpx.ConnectError("offline before acknowledgement", request=request)
            return httpx.Response(200, json={"status": "saved"})
        if request.url.path.endswith("/jobs"):
            return httpx.Response(200, json={"jobs": []})
        return httpx.Response(404)
    client = httpx.Client(base_url="https://backend.example", transport=httpx.MockTransport(handler))
    service = Companion("https://backend.example", "x" * 40, tmp_path, client=client)
    service.manifest.put(job, "claimed")
    service._handle(job)
    assert service.target(job).exists()
    assert len(service.manifest.pending()) == 1
    service.close()
    next_client = httpx.Client(base_url="https://backend.example", transport=httpx.MockTransport(handler))
    restarted = Companion("https://backend.example", "x" * 40, tmp_path, client=next_client)
    restarted.poll_once()
    assert counts == {"download": 1, "ack": 2}
    assert restarted.manifest.pending() == []
    restarted.close()


def test_invalid_workbook_is_not_acknowledged(tmp_path, payload):
    job, _ = companion_job(payload)
    content = b"not a workbook"
    job.update(checksum=hashlib.sha256(content).hexdigest(), size=len(content))
    acks = []
    def handler(request):
        if request.url.path.endswith("/ack"):
            acks.append(True)
        return httpx.Response(200, content=content)
    client = httpx.Client(base_url="https://backend.example", transport=httpx.MockTransport(handler))
    service = Companion("https://backend.example", "x" * 40, tmp_path, client=client)
    service._handle(job)
    assert acks == [] and not service.target(job).exists()
    service.close()


def test_companion_rejects_path_escape(tmp_path, payload):
    job, _ = companion_job(payload)
    service = Companion("https://backend.example", "x" * 40, tmp_path,
                        client=httpx.Client(base_url="https://backend.example", transport=httpx.MockTransport(lambda r: httpx.Response(200))))
    job["relative_path"] = "../../NIFTY_trade_001.xlsx"
    with pytest.raises(ValueError, match="Unsafe"):
        service.target(job)
    service.close()


def test_companion_url_requires_tls_except_loopback():
    assert validate_backend_url("http://127.0.0.1:8000") == "http://127.0.0.1:8000"
    with pytest.raises(ValueError, match="HTTPS"):
        validate_backend_url("http://remote.example")
    with pytest.raises(ValueError):
        validate_backend_url("https://token@remote.example")
