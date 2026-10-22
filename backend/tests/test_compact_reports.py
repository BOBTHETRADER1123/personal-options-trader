import hashlib
import io
from contextlib import closing

from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app import reporting
from companion.companion import verify_file


def trade():
    return {"id": "compact_test", "broker": "dhan", "name": "Paper call", "underlying": "NIFTY",
            "status": "closed", "created_at": "2026-10-21T04:00:00+00:00", "closed_at": "2026-10-21T09:00:00+00:00",
            "gross_pnl": 1050, "charges": 50, "net_pnl": 1000, "mfe": 2000, "mae": -1500,
            "mfe_at": "2026-10-21T07:00:00+00:00", "mae_at": "2026-10-21T05:00:00+00:00",
            "demo": True, "legs": [{"symbol": "NIFTY_TEST", "expiry": "2026-10-27", "strike": 23000,
                "option_type": "CE", "side": "BUY", "lots": 1, "lot_size": 65, "entry_price": 100, "exit_price": 116.15}]}


def test_compact_report_preserves_observed_extrema_without_ticks(tmp_path):
    content = reporting.build_workbook({"trade": trade(), "quotes": [{"ltp": 999}] * 1000, "timeline": [{"net_pnl": 123}] * 1000})
    with closing(load_workbook(io.BytesIO(content))) as book:
        assert set(book.sheetnames) == reporting.SUMMARY_SHEETS
        values = dict(book["Trade Summary"].iter_rows(max_col=2, values_only=True))
        assert values["Maximum profit reached (INR)"] == 2000
        assert values["Maximum loss reached (INR)"] == 1500
        assert values["Net closing P&L (INR)"] == 1000
        assert values["Maximum profit time (IST)"].hour == 12
        assert values["Maximum profit time (IST)"].minute == 30
        assert values["Maximum loss time (IST)"].hour == 10
    path = tmp_path / "compact.xlsx"
    path.write_bytes(content)
    verify_file(path, hashlib.sha256(content).hexdigest(), len(content))
    assert len(content) < 12000


def test_all_losing_trade_does_not_invent_profit_time():
    value = trade()
    value.update(mfe=-20, mae=-1500)
    with closing(load_workbook(io.BytesIO(reporting.build_workbook({"trade": value})))) as book:
        rows = dict(book["Trade Summary"].iter_rows(max_col=2, values_only=True))
        assert rows["Maximum profit reached (INR)"] == 0
        assert rows["Maximum profit time (IST)"] is None
        assert rows["Highest observed net P&L / MFE (INR)"] == -20


def test_queue_and_metadata_do_not_copy_or_read_raw_history(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'jobs.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    with factory() as db:
        job = reporting.enqueue_export(db, trade(), [{"private_raw": "x" * 10000}], [], [{"ltp": 100}])
        db.commit()
        assert set(job.payload) == {"format", "trade"}
        job.file_bytes = b"private_xlsx_bytes"
        db.commit()
    statements = []
    event.listen(engine, "before_cursor_execute", lambda conn, cursor, statement, params, context, many: statements.append(statement))
    with factory() as db:
        rows = reporting.job_metadata(db)
        assert rows[0]["demo"] is True
        assert rows[0]["trade_id"] == "compact_test"
        assert "file_bytes" not in statements[0]
        assert ', export_jobs.payload,' not in statements[0]
    engine.dispose()
