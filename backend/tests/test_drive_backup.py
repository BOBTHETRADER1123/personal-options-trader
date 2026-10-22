"""Verified Drive outbox checks; no Google account or cloud database is used."""
import hashlib
import importlib.util
import io
import json
import sqlite3
from contextlib import closing
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app import reporting


spec = importlib.util.spec_from_file_location('drive_backup_tests', Path(__file__).resolve().parents[2] / 'scripts' / 'drive_backup.py')
backup = importlib.util.module_from_spec(spec)
spec.loader.exec_module(backup)


@pytest.fixture
def trade():
    return {'id': 'test-trade', 'broker': 'dhan', 'underlying': 'NIFTY', 'name': 'Sample call',
            'status': 'closed', 'demo': True, 'created_at': '2026-10-22T04:00:00+00:00',
            'closed_at': '2026-10-22T08:00:00+00:00', 'net_pnl': 80, 'gross_pnl': 100,
            'charges': 20, 'mfe': 120, 'mae': -30, 'mfe_at': '2026-10-22T07:00:00+00:00',
            'mae_at': '2026-10-22T05:00:00+00:00', 'private_token': 'NEVER_EXPORT_TEST_TOKEN',
            'legs': [{'symbol': 'TEST_CALL', 'strike': 24000, 'expiry': '2026-10-27',
                      'option_type': 'CE', 'side': 'BUY', 'lots': 1, 'lot_size': 50,
                      'entry_price': 100, 'exit_price': 102, 'session_token': 'NEVER_EXPORT_LEG_TOKEN'}]}


@pytest.fixture
def project(tmp_path, trade):
    root = tmp_path / 'paper-app'
    data = root / 'OptionsTraderData'
    data.mkdir(parents=True)
    (root / 'companion').mkdir()
    (root / '.env').write_text('DATABASE_URL=postgresql://secret:password@not-contacted.invalid/paper\n', encoding='utf-8')
    with sqlite3.connect(data / 'trader.db') as connection:
        connection.execute('CREATE TABLE trades(id TEXT PRIMARY KEY,broker TEXT,status TEXT,data TEXT)')
        connection.execute('INSERT INTO trades VALUES (?,?,?,?)', (trade['id'], trade['broker'], 'closed', json.dumps(trade)))
        opened = {**trade, 'id': 'still-open', 'status': 'open'}
        connection.execute('INSERT INTO trades VALUES (?,?,?,?)', ('still-open', 'dhan', 'open', json.dumps(opened)))
    return root


def prepare(root):
    return backup.prepare(root, report_builder=reporting.build_workbook, verifier=reporting.verify_workbook)


def queue(root):
    return backup.BackupQueue(root / '.desktop' / 'drive-backup')


def test_prepares_only_closed_compact_summary_without_modifying_database(project):
    database = project / 'OptionsTraderData' / 'trader.db'
    original = database.read_bytes()
    result = prepare(project)
    assert result['pending_count'] == 1
    assert result['issues'] == []
    assert result['source_records_deleted'] is False
    item = result['pending'][0]
    assert set(item) == {'key', 'filename', 'source', 'size', 'md5', 'sha256', 'path', 'folder_id'}
    assert 'password' not in json.dumps(result) and 'NEVER_EXPORT' not in json.dumps(result)
    content = Path(item['path']).read_bytes()
    assert hashlib.md5(content).hexdigest() == item['md5']
    with closing(load_workbook(io.BytesIO(content))) as book:
        assert set(book.sheetnames) == {'Trade Summary', 'Strategy Legs'}
        cells = [row for sheet in book for row in sheet.iter_rows(values_only=True)]
        assert 'NEVER_EXPORT' not in str(cells)
        summary = dict(book['Trade Summary'].iter_rows(max_col=2, values_only=True))
        assert summary['Maximum profit reached (INR)'] == 120
        assert summary['Maximum loss reached (INR)'] == 30
        assert summary['Maximum profit time (IST)'].hour == 12
    assert database.read_bytes() == original


def test_prepare_and_confirm_are_idempotent(project):
    first = prepare(project)['pending'][0]
    content = Path(first['path']).read_bytes()
    second = prepare(project)['pending'][0]
    assert first == second
    assert Path(second['path']).read_bytes() == content
    archive = queue(project)
    try:
        assert archive.confirm(first['key'], 'drive_file_12345', first['md5'], first['size'])['status'] == 'uploaded'
        archive.confirm(first['key'], 'drive_file_12345', first['md5'], first['size'])
        with pytest.raises(ValueError, match='another Drive file'):
            archive.confirm(first['key'], 'another_drive_file', first['md5'], first['size'])
    finally:
        archive.close()
    result = prepare(project)
    assert result['pending_count'] == 0 and result['uploaded_count'] == 1


@pytest.mark.parametrize('change', ['md5', 'size'])
def test_rejects_remote_verification_mismatch_and_keeps_pending(project, change):
    item = prepare(project)['pending'][0]
    archive = queue(project)
    try:
        with pytest.raises(ValueError, match='does not match'):
            archive.confirm(item['key'], 'drive_file_12345', '0' * 32 if change == 'md5' else item['md5'],
                            item['size'] + (change == 'size'))
        assert len(archive.pending()) == 1
    finally:
        archive.close()


def test_changed_outbox_cannot_be_confirmed_and_prepare_repairs_it(project):
    item = prepare(project)['pending'][0]
    Path(item['path']).write_bytes(b'corrupt copy')
    archive = queue(project)
    try:
        with pytest.raises(ValueError, match='has changed'):
            archive.confirm(item['key'], 'drive_file_12345', item['md5'], item['size'])
        with pytest.raises(ValueError, match='changed'):
            archive.pending()
    finally:
        archive.close()
    repaired = prepare(project)['pending'][0]
    assert repaired['key'] == item['key']
    reporting.verify_workbook(Path(repaired['path']).read_bytes())


def test_changed_closed_summary_creates_new_version(project):
    first = prepare(project)['pending'][0]
    with sqlite3.connect(project / 'OptionsTraderData' / 'trader.db') as connection:
        value = json.loads(connection.execute('SELECT data FROM trades WHERE id=?', ('test-trade',)).fetchone()[0])
        value['mfe'] = 150
        connection.execute('UPDATE trades SET data=? WHERE id=?', (json.dumps(value), 'test-trade'))
    result = prepare(project)
    assert result['pending_count'] == 2
    assert len({item['key'] for item in result['pending']}) == 2
    assert first['key'] in {item['key'] for item in result['pending']}


def add_companion(root, folder, content, path=None, status='acknowledged', checksum=None):
    folder.mkdir(parents=True, exist_ok=True)
    path = path or folder / 'saved.xlsx'
    path.write_bytes(content)
    metadata = folder / '.companion'
    metadata.mkdir(exist_ok=True)
    with sqlite3.connect(metadata / 'manifest.sqlite3') as connection:
        connection.execute('CREATE TABLE jobs(id TEXT PRIMARY KEY,data TEXT,status TEXT,error TEXT)')
        job = {'id': 'test-trade', 'checksum': checksum or hashlib.sha256(content).hexdigest(),
               'size': len(content), 'local_path': str(path), 'filename': path.name, 'claim_token': 'PRIVATE_CLAIM_TOKEN'}
        connection.execute('INSERT INTO jobs VALUES (?,?,?,NULL)', ('test-trade', json.dumps(job), status))
    (root / 'companion' / '.env.online').write_text(f'TRADES_FOLDER={folder.as_posix()}\nCOMPANION_TOKEN=PRIVATE_COMPANION_TOKEN\n', encoding='utf-8')


def test_preserves_verified_legacy_workbook_bytes(project, trade, tmp_path):
    content = reporting.build_workbook({'format': 'detailed', 'trade': trade, 'events': [], 'timeline': [], 'quotes': []})
    add_companion(project, tmp_path / 'online-reports', content)
    result = prepare(project)
    legacy = next(item for item in result['pending'] if item['source'] == 'verified_report')
    assert Path(legacy['path']).read_bytes() == content
    assert result['pending_count'] == 2  # Same trade ID in local/cloud cannot collide.
    assert 'PRIVATE' not in json.dumps(result)


@pytest.mark.parametrize('case', ['outside', 'checksum', 'unverified', 'invalid_xlsx'])
def test_companion_receipt_and_folder_boundaries_are_required(project, trade, tmp_path, case):
    content = reporting.build_workbook({'format': 'summary', 'trade': trade}) if case != 'invalid_xlsx' else b'not Excel'
    folder = tmp_path / 'online-reports'
    add_companion(project, folder, content,
                  path=tmp_path / 'unrelated.xlsx' if case == 'outside' else None,
                  status='claimed' if case == 'unverified' else 'acknowledged',
                  checksum='0' * 64 if case == 'checksum' else None)
    result = prepare(project)
    assert result['pending_count'] == 1
    assert all(item['source'] == 'local_summary' for item in result['pending'])
    assert bool(result['issues']) is (case != 'unverified')


def test_local_database_is_opened_readonly(project):
    connection = backup.readonly(project / 'OptionsTraderData' / 'trader.db')
    try:
        with pytest.raises(sqlite3.OperationalError, match='readonly'):
            connection.execute('DELETE FROM trades')
    finally:
        connection.close()


def test_formatter_import_uses_memory_database_and_restores_environment(tmp_path, monkeypatch):
    monkeypatch.setenv('DATABASE_URL', 'postgresql://private:private@not-contacted.invalid/paper')
    monkeypatch.setenv('DATA_DIR', 'old-data-value')
    calls = []
    def fake_import(name):
        calls.append((name, backup.os.environ['DATABASE_URL'], backup.os.environ['DATA_DIR']))
        return SimpleNamespace(build_workbook=lambda value: b'x', verify_workbook=lambda value: None)
    monkeypatch.setattr(backup.importlib, 'import_module', fake_import)
    backup.reporting_functions(tmp_path)
    assert calls == [('app.reporting', 'sqlite:///:memory:', str(tmp_path))]
    assert backup.os.environ['DATABASE_URL'].startswith('postgresql://private:')
    assert backup.os.environ['DATA_DIR'] == 'old-data-value'


def test_cli_failure_does_not_print_private_exception(capsys, monkeypatch):
    monkeypatch.setattr(backup, 'prepare', lambda: (_ for _ in ()).throw(RuntimeError('SECRET_PASSWORD')))
    assert backup.main(['prepare']) == 1
    result = capsys.readouterr().out
    assert 'SECRET' not in result
    assert 'Source records were retained' in json.loads(result)['error']
